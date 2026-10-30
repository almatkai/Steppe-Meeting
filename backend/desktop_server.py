"""
Steppe Meeting Desktop - Local Backend Server
Autonomous local backend engine with local SQLite, local Whisper STT,
and local Ollama / OpenAI-compatible LLM endpoints.
"""

import asyncio
from contextlib import asynccontextmanager, contextmanager
import json
import io
import logging
import os
import tempfile

# Ensure temp directory exists and is writable (prevents SpooledTemporaryFile failure on >1MB uploads)
current_tmp = os.environ.get("TMPDIR")
if current_tmp and not os.path.isdir(current_tmp):
    del os.environ["TMPDIR"]
    tempfile.tempdir = None

try:
    os.makedirs(tempfile.gettempdir(), exist_ok=True)
except Exception:
    tempfile.tempdir = "/tmp"
import re
import sqlite3
import subprocess
import sys
import time
import uuid
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional
import urllib.parse

import httpx
import uvicorn
from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, Query, Request, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from pydantic import BaseModel, ConfigDict, Field

# Add current backend folder to python path
BASE_DIR = Path(__file__).resolve().parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

# Keep all writable runtime data outside the source tree / application bundle.
# Existing backend/data contents are copied once for backward compatibility.
from app_data import prepare_app_data_dir

LEGACY_DATA_DIR = BASE_DIR / "data"
DATA_DIR = prepare_app_data_dir(LEGACY_DATA_DIR)
AUDIO_DIR = DATA_DIR / "audio"
EXPORTS_DIR = DATA_DIR / "exports"
TEMPLATES_DIR = DATA_DIR / "templates"
DB_PATH = DATA_DIR / "meetings.db"

# Import local engine modules
from engine.config import settings
from engine.generation_route import GenerationRoute, use_route
from live_captions import create_router as create_captions_router
from recording_streams import RecordingStreams
from recording_cache import retain_recording
from generation_jobs import GenerationJobs
from release_updates import checker as update_checker
from update_service import UpdateService
import secrets
from engine import cli_llm, chunking, docx_template, docx_preview, template_generation, template_profile, template_protocol, pdf_export, llm, prompts, protocol, summary, vector_store, embeddings, indexer, retriever, whisper_service

# LLM telemetry is runtime data too, not a repo-relative directory.
settings.LLM_LOG_DIR = str(DATA_DIR / "logs" / "llm")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("steppe.desktop")

# ---------------------------------------------------------------------------
# Database Management (SQLite)
# ---------------------------------------------------------------------------

@contextmanager
def get_db():
    conn = sqlite3.connect(str(DB_PATH), check_same_thread=False, timeout=15.0)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
    finally:
        conn.close()

def init_db():
    with get_db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS meetings (
                id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                created_at TEXT NOT NULL,
                status TEXT NOT NULL,
                duration_seconds REAL DEFAULT 0,
                source_language TEXT DEFAULT 'multi',
                agenda TEXT DEFAULT '',
                participants TEXT DEFAULT '[]',
                audio_filename TEXT DEFAULT '',
                audio_path TEXT DEFAULT '',
                transcript_text TEXT DEFAULT '',
                transcript_segments TEXT DEFAULT '[]',
                protocol_ru TEXT DEFAULT '{}',
                protocol_kz TEXT DEFAULT '{}',
                summary_ru TEXT DEFAULT '{}',
                summary_kz TEXT DEFAULT '{}',
                error_message TEXT DEFAULT ''
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS chat_sessions (
                id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS chat_messages (
                id TEXT PRIMARY KEY,
                session_id TEXT NOT NULL,
                role TEXT NOT NULL,
                content TEXT NOT NULL,
                citations TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (session_id) REFERENCES chat_sessions(id) ON DELETE CASCADE
            )
        """)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_chat_messages_session ON chat_messages(session_id)")

        # Protocol Templates table
        conn.execute("""
            CREATE TABLE IF NOT EXISTS protocol_templates (
                id TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                description TEXT DEFAULT '',
                additional_prompt TEXT DEFAULT '',
                detail_level TEXT DEFAULT 'concise',
                status TEXT DEFAULT 'approved',
                docx_path TEXT NOT NULL,
                source_docx_path TEXT NOT NULL,
                slots TEXT NOT NULL,
                schema_json TEXT NOT NULL,
                style_config TEXT NOT NULL,
                template_profile TEXT DEFAULT '{}',
                render_ready INTEGER DEFAULT 0,
                test_values TEXT DEFAULT '{}',
                test_docx_path TEXT DEFAULT '',
                test_pdf_path TEXT DEFAULT '',
                created_at TEXT NOT NULL
            )
        """)

        # Upgrade protocol_templates for test_pdf_path
        try:
            conn.execute("ALTER TABLE protocol_templates ADD COLUMN test_pdf_path TEXT DEFAULT ''")
        except sqlite3.OperationalError:
            pass

        # Upgrade meetings table for templates
        for col, col_type in [
            ("template_id", "TEXT DEFAULT ''"),
            ("template_values_ru", "TEXT DEFAULT '{}'"),
            ("template_values_kz", "TEXT DEFAULT '{}'"),
            ("generation_attempt", "TEXT DEFAULT '{}'"),
        ]:
            try:
                conn.execute(f"ALTER TABLE meetings ADD COLUMN {col} {col_type}")
            except sqlite3.OperationalError:
                pass

        conn.commit()
        vector_store.ensure_schema(conn)


def migrate_stored_paths() -> None:
    """Point absolute paths in a migrated database at the new app data root."""
    legacy_prefix = str(LEGACY_DATA_DIR.resolve())
    data_prefix = str(DATA_DIR.resolve())
    if legacy_prefix == data_prefix:
        return

    with get_db() as conn:
        conn.execute(
            "UPDATE meetings SET audio_path = REPLACE(audio_path, ?, ?) "
            "WHERE audio_path LIKE ?",
            (legacy_prefix, data_prefix, f"{legacy_prefix}%"),
        )
        for column in ("docx_path", "source_docx_path", "test_docx_path"):
            conn.execute(
                f"UPDATE protocol_templates SET {column} = REPLACE({column}, ?, ?) "
                f"WHERE {column} LIKE ?",
                (legacy_prefix, data_prefix, f"{legacy_prefix}%"),
            )
        conn.commit()


init_db()
migrate_stored_paths()

def seed_default_template():
    with get_db() as conn:
        row = conn.execute(
            "SELECT 1 FROM protocol_templates WHERE id = ?",
            ("default-protocol-template",),
        ).fetchone()
        if row:
            return

        candidate_paths = [
            BASE_DIR.parent / "public" / "templates" / "ideal-protocol-template.docx",
            BASE_DIR.parent.parent / "jynalys" / "jinalys-new-admin-frontend" / "public" / "templates" / "ideal-protocol-template.docx",
        ]
        sample_path = None
        for p in candidate_paths:
            if p.exists():
                sample_path = p
                break

        if not sample_path:
            return

        try:
            with open(sample_path, "rb") as f:
                docx_bytes = f.read()

            source_parsed = docx_template.parse_docx_template(docx_bytes)
            working_bytes = docx_template.normalize_visual_markers(docx_bytes)
            parsed = docx_template.parse_docx_template(working_bytes)

            source_slots = {slot.key: slot for slot in source_parsed.slots}
            loop_iters = docx_template._loop_iterables(working_bytes)
            for slot in parsed.slots:
                source_slot = source_slots.get(slot.key)
                if source_slot and source_slot.source != "placeholder":
                    slot.label = source_slot.label
                    slot.value_type = source_slot.value_type
                    slot.repeat = source_slot.repeat
                    slot.omit_when_empty = source_slot.omit_when_empty
                elif slot.key in loop_iters:
                    slot.value_type = "list[object]"

            name = "Типовой протокол (Стандарт РК)"
            descriptor = docx_template.parsed_to_descriptor(parsed, name)
            template_id = "default-protocol-template"
            target_working = TEMPLATES_DIR / f"{template_id}.docx"
            target_source = TEMPLATES_DIR / f"{template_id}_source.docx"

            with open(target_working, "wb") as f:
                f.write(working_bytes)
            with open(target_source, "wb") as f:
                f.write(docx_bytes)

            now = datetime.now(timezone.utc).isoformat()
            conn.execute("""
                INSERT OR REPLACE INTO protocol_templates (
                    id, name, description, additional_prompt, detail_level, status,
                    docx_path, source_docx_path, slots, schema_json, style_config,
                    template_profile, render_ready, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                template_id,
                name,
                "Идеальный образец шаблона протокола совещания с таблицами, списком решений и участниками.",
                "",
                "concise",
                "approved",
                str(target_working),
                str(target_source),
                json.dumps(descriptor["slots"], ensure_ascii=False),
                json.dumps(descriptor["schema_json"], ensure_ascii=False),
                json.dumps(descriptor["style_config"], ensure_ascii=False),
                json.dumps({}, ensure_ascii=False),
                1 if descriptor["render_ready"] else 0,
                now,
            ))
            conn.commit()
            logger.info("Default protocol template seeded: %s", template_id)
        except Exception as e:
            logger.warning("Failed to seed default template: %s", e)

seed_default_template()

def get_setting(key: str, default: str = "") -> str:
    with get_db() as conn:
        row = conn.execute("SELECT value FROM settings WHERE key = ?", (key,)).fetchone()
        return row["value"] if row else default

def set_setting(key: str, value: str):
    with get_db() as conn:
        conn.execute("INSERT OR REPLACE INTO settings (key, value) VALUES (?, ?)", (key, value))
        conn.commit()

# Default Settings
DEFAULT_LLM_URL = os.environ.get("LLM_BASE_URL") or get_setting("llm_base_url", "http://localhost:11434/v1")
DEFAULT_LLM_MODEL = os.environ.get("LLM_MODEL") or get_setting("llm_model", "qwen2.5:latest")
DEFAULT_LLM_KEY = os.environ.get("LLM_API_KEY") or get_setting("llm_api_key", "")
DEFAULT_WHISPER_MODE = os.environ.get("WHISPER_MODE") or get_setting("whisper_mode", "local")
DEFAULT_WHISPER_LOCAL_MODEL = os.environ.get("WHISPER_LOCAL_MODEL") or get_setting("whisper_local_model", "small")
DEFAULT_WHISPER_DEVICE = os.environ.get("WHISPER_DEVICE") or get_setting("whisper_device", "auto")
DEFAULT_WHISPER_URL = os.environ.get("WHISPER_BASE_URL") or get_setting("whisper_base_url", "http://localhost:8000/v1")
DEFAULT_WHISPER_MODEL = os.environ.get("WHISPER_MODEL") or get_setting("whisper_model", "whisper-1")
DEFAULT_WHISPER_KEY = os.environ.get("WHISPER_API_KEY") or get_setting("whisper_api_key", "")

set_setting("llm_base_url", DEFAULT_LLM_URL)
set_setting("llm_model", DEFAULT_LLM_MODEL)
set_setting("llm_api_key", DEFAULT_LLM_KEY)
set_setting("whisper_mode", DEFAULT_WHISPER_MODE)
set_setting("whisper_local_model", DEFAULT_WHISPER_LOCAL_MODEL)
set_setting("whisper_device", DEFAULT_WHISPER_DEVICE)
set_setting("whisper_base_url", DEFAULT_WHISPER_URL)
set_setting("whisper_model", DEFAULT_WHISPER_MODEL)
set_setting("whisper_api_key", DEFAULT_WHISPER_KEY)

def get_models_url(base_url: str) -> str:
    """Normalizes any provider base URL to the /models endpoint."""
    url = (base_url or "").strip().rstrip("/")
    if url.endswith("/models"):
        return url
    if url.endswith("/v1"):
        return f"{url}/models"
    return f"{url}/v1/models"

def get_ollama_tags_url(base_url: str) -> str:
    url = (base_url or "").strip().rstrip("/")
    if url.endswith("/v1"):
        url = url[:-3]
    return f"{url}/api/tags"

def get_auth_headers(api_key: Optional[str]) -> Dict[str, str]:
    if not api_key or not api_key.strip():
        return {}
    key = api_key.strip()
    if key.lower().startswith("bearer "):
        key = key[7:].strip()
    return {"Authorization": f"Bearer {key}"}

def get_cli_config():
    try:
        value = json.loads(get_setting("llm_cli", "{}"))
        return value if isinstance(value, dict) else {}
    except ValueError:
        return {}


def sync_generation_settings():
    settings.LLM_MODE = get_setting("llm_mode", "api")
    settings.LLM_CLI_PROVIDER = get_setting("llm_cli_provider", "opencode")
    settings.LLM_CLI = get_cli_config()
    llm_url = get_setting("llm_base_url", "http://localhost:11434/v1")
    llm_model = get_setting("llm_model", "qwen2.5:latest")
    llm_api_key = get_setting("llm_api_key", "")
    settings.LLM_BASE_URL = llm_url
    settings.LLM_MODEL = llm_model
    settings.LLM_API_KEY = llm_api_key
    if "api.openai.com" in llm_url:
        settings.LLM_PRIMARY = "openai"
        settings.OPENAI_API_KEY = llm_api_key
        settings.OPENAI_LLM_MODEL = llm_model
    else:
        settings.LLM_PRIMARY = "local"
    settings.STT_SERVICE_BASE_URL = get_setting("whisper_base_url", "http://localhost:8000/v1")

sync_generation_settings()

# ---------------------------------------------------------------------------
# FastAPI Application
# ---------------------------------------------------------------------------

def write_generation_status(meeting_id, status, error):
    with get_db() as conn:
        conn.execute("UPDATE meetings SET status=?, error_message=? WHERE id=?", (status, error, meeting_id))
        conn.commit()


generation_jobs = GenerationJobs(write_generation_status)


def recover_orphaned_generations():
    # No owned task survives process restart; never pretend it is still running.
    with get_db() as conn:
        conn.execute("UPDATE meetings SET status='cancelled', error_message='' WHERE status IN ('generating', 'cancelling')")
        conn.execute("UPDATE meetings SET status='error', error_message=? WHERE status='transcribing'",
                     ("Распознавание прервано перезапуском приложения. Аудио сохранено; запустите обработку заново.",))
        conn.commit()


@asynccontextmanager
async def lifespan(app):
    recover_orphaned_generations()
    try:
        yield
    finally:
        await generation_jobs.shutdown()


app = FastAPI(
    lifespan=lifespan,
    title="Meeting Protocol Desktop API",
    version="1.0.0",
    description="Local backend engine for meeting protocol and summary generation",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.exception_handler(HTTPException)
async def http_exception_handler(request: Request, exc: HTTPException):
    if exc.status_code == 400 and exc.detail == "There was an error parsing the body":
        cause = getattr(exc, "__cause__", None)
        logger.error(
            "Body parsing failed on %s %s: cause=%r",
            request.method, request.url.path, cause, exc_info=cause or exc
        )
    return JSONResponse(status_code=exc.status_code, content={"detail": exc.detail})


recording_streams = RecordingStreams(get_db, lambda key, default=None: get_setting(key, default))
app.include_router(create_captions_router(get_setting, recording_streams))


def update_busy():
    reasons = []
    if any(generation_jobs.active(key) for key in generation_jobs.jobs):
        reasons.append("generation")
    with get_db() as db:
        if db.execute("SELECT 1 FROM meetings WHERE status IN ('transcribing','generating') LIMIT 1").fetchone():
            reasons.append("processing")
    # Durable transcript checkpoints are not capture ownership. Browser capture
    # is fenced by native webview ownership; chunk work owns these process-local
    # locks (including requests waiting for an earlier chunk). Restart orphans
    # must remain recoverable without impersonating a live microphone.
    if any(lock.locked() for lock in recording_streams.locks.values()):
        reasons.append("recording")
    if update_active_requests:
        reasons.append("request")
    return reasons


update_active_requests = 0
updater = UpdateService(DATA_DIR / "updates", update_checker, os.environ.get("STEPPE_APP_VERSION", "development"),
                        lambda: get_update_channel(), update_busy)


@app.middleware("http")
async def update_work_fence(request: Request, call_next):
    global update_active_requests
    path = request.url.path
    work = request.method != "GET" and not path.startswith("/api/v1/system/updates")
    if work and updater.installing:
        return JSONResponse(status_code=409, content={"detail": "Приложение готовится к обновлению. Дождитесь перезапуска."})
    if work:
        update_active_requests += 1
        updater.activity()
    try:
        response = await call_next(request)
    except BaseException:
        if work:
            update_active_requests -= 1
        raise
    if work:
        iterator = response.body_iterator
        async def counted_body():
            global update_active_requests
            try:
                async for chunk in iterator:
                    yield chunk
            finally:
                update_active_requests -= 1
        response.body_iterator = counted_body()
    return response


class UpdateSelectionRequest(BaseModel):
    selection_id: str = Field(max_length=64)


class UpdateJobRequest(BaseModel):
    job_id: str = Field(max_length=64)


@app.get("/api/v1/system/updates/job")
async def update_job_status():
    return updater.state()


@app.post("/api/v1/system/updates/download")
async def update_download(req: UpdateSelectionRequest):
    try:
        return await updater.begin(req.selection_id)
    except ValueError as error:
        return JSONResponse(status_code=409, content={"detail": str(error), "error_code": "invalid_update_request"})


@app.post("/api/v1/system/updates/cancel")
async def update_cancel(req: UpdateJobRequest):
    try:
        return await updater.cancel(req.job_id)
    except ValueError as error:
        return JSONResponse(status_code=409, content={"detail": str(error), "error_code": "invalid_update_request"})


@app.post("/api/v1/system/updates/native-install")
async def update_native_install(req: UpdateJobRequest, request: Request):
    token = os.environ.get("STEPPE_NATIVE_UPDATE_TOKEN", "")
    if not token or not secrets.compare_digest(request.headers.get("x-steppe-native", ""), token):
        raise HTTPException(403, "Установка доступна только из приложения.")
    try:
        return await updater.install(req.job_id)
    except ValueError as error:
        return JSONResponse(status_code=409, content={"detail": str(error), "error_code": "invalid_update_request"})


# ---------------------------------------------------------------------------
# Schemas
# ---------------------------------------------------------------------------

class CreateMeetingRequest(BaseModel):
    title: str
    source_language: str = "multi"
    agenda: Optional[str] = ""
    participants: Optional[List[Dict[str, Any]]] = []
    transcript_text: Optional[str] = ""
    template_id: Optional[str] = None

class UpdateMeetingRequest(BaseModel):
    title: Optional[str] = None
    source_language: Optional[str] = None
    agenda: Optional[str] = None
    participants: Optional[List[Dict[str, Any]]] = None
    transcript_text: Optional[str] = None
    protocol_ru: Optional[Dict[str, Any]] = None
    protocol_kz: Optional[Dict[str, Any]] = None
    summary_ru: Optional[Dict[str, Any]] = None
    summary_kz: Optional[Dict[str, Any]] = None
    template_id: Optional[str] = None

class TemplateTestRequest(BaseModel):
    transcript: Optional[str] = None
    detail_level: Optional[str] = None

class CliOptions(BaseModel):
    path: Optional[str] = None
    model: Optional[str] = None


class CliRequest(BaseModel):
    provider: str
    path: Optional[str] = None


class CliTestRequest(CliRequest):
    model: Optional[str] = None


class GenerateMeetingRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    mode: Literal["api", "cli"]
    provider: Optional[Literal["api", "opencode", "codex", "claude", "antigravity"]] = None
    model: Optional[str] = Field(default=None, max_length=512)


def generation_config():
    """Read one coherent saved configuration without mutating shared settings."""
    with get_db() as conn:
        values = dict(conn.execute("SELECT key,value FROM settings WHERE key IN ('llm_mode','llm_cli_provider','llm_cli','llm_base_url','llm_model','llm_api_key')"))
    try:
        cli = json.loads(values.get("llm_cli", "{}"))
    except ValueError:
        cli = {}
    return {
        "mode": values.get("llm_mode", "api"),
        "provider": values.get("llm_cli_provider", "opencode"),
        "cli": cli if isinstance(cli, dict) else {},
        "base_url": values.get("llm_base_url", DEFAULT_LLM_URL),
        "model": values.get("llm_model", DEFAULT_LLM_MODEL),
        "api_key": values.get("llm_api_key", DEFAULT_LLM_KEY),
    }


def resolve_generation_route(request: GenerateMeetingRequest | None = None):
    config = generation_config()
    mode = request.mode if request else config["mode"]
    provider = "api"
    if mode == "cli":
        provider = request.provider if request and request.provider else config["provider"]
    if request and mode == "api" and request.provider not in (None, "api"):
        raise ValueError("Для режима API выберите настроенный API-провайдер")
    options = config["cli"].get(provider, {})
    if not isinstance(options, dict):
        options = {}
    model = request.model if request and request.model is not None else options.get("model", "") if mode == "cli" else config["model"]
    if not isinstance(model, str):
        raise ValueError("Некорректное имя модели")
    model = model.strip()
    if any(character in model for character in ('\x00', '\n', '\r')) or len(model) > 512:
        raise ValueError("Некорректное имя модели")
    if mode == "cli":
        path = options.get("path", "") or ""
        cli_llm.validate(provider, path, model, require_model=True)
        return GenerationRoute(mode, provider, model, path=path, timeout=settings.LLM_TIMEOUT)
    if mode != "api":
        raise ValueError("Неизвестный способ генерации")
    parsed = urllib.parse.urlparse(config["base_url"])
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise ValueError("Укажите HTTP или HTTPS URL провайдера в Настройках AI")
    if not model:
        raise ValueError("Укажите модель для генерации")
    return GenerationRoute(mode, "api", model, base_url=config["base_url"], api_key=config["api_key"], timeout=settings.LLM_TIMEOUT)


def start_meeting_generation(meeting_id, route=None):
    route = route or resolve_generation_route()
    if generation_jobs.active(meeting_id):
        raise ValueError("Generation already running")
    attempt = {**route.identity(), "id": str(uuid.uuid4()), "started_at": datetime.now(timezone.utc).isoformat()}
    with get_db() as conn:
        conn.execute("UPDATE meetings SET generation_attempt=? WHERE id=?", (json.dumps(attempt), meeting_id))
        conn.commit()

    async def generate(id):
        with use_route(route):
            await run_recording_pipeline(id)
    generation_jobs.start(meeting_id, generate)
    return attempt


def start_saved_meeting_generation(meeting_id):
    """Auto handoff keeps saved transcription usable when settings need repair."""
    try:
        start_meeting_generation(meeting_id)
        return "generating"
    except ValueError as error:
        write_generation_status(meeting_id, "error", "Не удалось начать генерацию: " + str(error))
        return "error"


class SystemConfigRequest(BaseModel):
    update_channel: Optional[Literal["stable", "beta"]] = None
    llm_mode: Optional[str] = None
    llm_cli_provider: Optional[str] = None
    llm_cli: Optional[Dict[str, CliOptions]] = None
    llm_base_url: Optional[str] = None
    llm_model: Optional[str] = None
    llm_api_key: Optional[str] = None
    whisper_mode: Optional[str] = None
    whisper_local_model: Optional[str] = None
    whisper_device: Optional[str] = None
    whisper_base_url: Optional[str] = None
    whisper_model: Optional[str] = None
    whisper_api_key: Optional[str] = None

class UpdateChannelRequest(BaseModel):
    channel: Literal["stable", "beta"]

class WhisperModelActionRequest(BaseModel):
    model: str

class ChatMessage(BaseModel):
    role: str
    content: str

class ChatRequest(BaseModel):
    meeting_id: Optional[str] = None
    messages: List[ChatMessage]
    context: Optional[Dict[str, Any]] = None

class EditRequest(BaseModel):
    selected_text: str
    instruction: str
    context: Optional[Dict[str, Any]] = None

# ---------------------------------------------------------------------------
# Meeting Endpoints
# ---------------------------------------------------------------------------

@app.post("/api/v1/meetings/create")
async def create_meeting(req: CreateMeetingRequest):
    meeting_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    status = "transcribed" if req.transcript_text.strip() else "draft"

    with get_db() as conn:
        conn.execute("""
            INSERT INTO meetings (
                id, title, created_at, status, source_language, agenda,
                participants, transcript_text, template_id
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            meeting_id,
            req.title or f"Совещание {datetime.now().strftime('%d.%m.%Y %H:%M')}",
            now,
            status,
            req.source_language,
            req.agenda or "",
            json.dumps(req.participants or [], ensure_ascii=False),
            req.transcript_text or "",
            req.template_id or "",
        ))
        conn.commit()

    return {"id": meeting_id, "status": status, "created_at": now}

@app.post("/api/v1/meetings/{meeting_id}/upload-audio")
async def upload_audio(meeting_id: str, file: UploadFile = File(...)):
    with get_db() as conn:
        row = conn.execute("SELECT id, source_language FROM meetings WHERE id = ?", (meeting_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Meeting not found")
        session = conn.execute("SELECT config FROM recording_streams WHERE meeting_id=?", (meeting_id,)).fetchone()

    clean_filename = Path(file.filename).name if file.filename else "recording.wav"
    # Remove any stray path separators
    clean_filename = clean_filename.replace("/", "_").replace("\\", "_")
    safe_filename = f"{meeting_id}_{clean_filename}"
    file_path = AUDIO_DIR / safe_filename

    mode = get_setting("whisper_mode", "local")
    config = json.loads(session["config"]) if session else {
        "mode": mode,
        "model": get_setting("whisper_local_model", "small") if mode == "local" else get_setting("whisper_model", "whisper-1"),
        "device": get_setting("whisper_device", "auto"),
        "language": row["source_language"],
    }
    config["streaming"] = bool(session)
    # Mini and wizard use this same endpoint. Publish only after stop/final data
    # has arrived and the durable regression copy is verified. Never acknowledge
    # a cache failure as a successful save or leave a partially overwritten file.
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=AUDIO_DIR, prefix=".upload-", delete=False) as output:
            temporary = Path(output.name)
            size = 0
            while content := await file.read(1024 * 1024):
                output.write(content)
                size += len(content)
            output.flush()
            os.fsync(output.fileno())
        cached = await asyncio.to_thread(retain_recording, temporary, meeting_id, config,
                                        root=DATA_DIR / "stt-regression", source_name=clean_filename)
        os.replace(temporary, file_path)
    except OSError as error:
        raise HTTPException(503, "Не удалось сохранить аудио и локальный кэш. Освободите место и повторите сохранение.") from error
    finally:
        if temporary and temporary.exists():
            temporary.unlink()

    with get_db() as conn:
        conn.execute("""
            UPDATE meetings 
            SET audio_filename = ?, audio_path = ?, status = 'draft'
            WHERE id = ?
        """, (clean_filename, str(file_path), meeting_id))
        conn.commit()

    return {"meeting_id": meeting_id, "filename": clean_filename, "size": size, "regression_case": cached["sha256"]}

class RecordingCompletionRequest(BaseModel):
    session_id: str
    chunk_count: int = Field(ge=0)
    duration_seconds: float = Field(ge=0, allow_inf_nan=False)


@app.post("/api/v1/meetings/{meeting_id}/process-recording")
async def process_recording(meeting_id: str, request: Optional[RecordingCompletionRequest] = None):
    """Finalize the saved streaming transcript, then generate without another STT pass."""
    with get_db() as conn:
        conn.execute("BEGIN IMMEDIATE")
        row = conn.execute("SELECT * FROM meetings WHERE id=?", (meeting_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Meeting not found")
        if request is None:
            # Installed clients may outlive a backend update. Never silently
            # re-transcribe their audio; allow retry only when text already exists.
            if row["status"] in ("generating", "cancelling", "cancelled", "completed", "error") and (row["transcript_text"] or "").strip():
                return {"meeting_id": meeting_id, "status": row["status"]}
            raise HTTPException(409, "Эта запись начата в старой версии Mini. Аудио сохранено в истории. Обновите приложение; запись можно восстановить из сохранённого аудио.")
        session = conn.execute("SELECT finalized FROM recording_streams WHERE id=? AND meeting_id=?",
                               (request.session_id, meeting_id)).fetchone()
        if not session:
            raise HTTPException(409, "Транскрипт этой записи не найден.")
        if session["finalized"] and (row["error_message"] or "").startswith("Запись сброшена."):
            raise HTTPException(409, "Запись сброшена. Начните новую запись в Mini.")
        if session["finalized"] == 1:
            return {"meeting_id": meeting_id, "status": row["status"]}
        if generation_jobs.active(meeting_id) or row["status"] == "transcribing":
            raise HTTPException(409, "Обработка совещания уже выполняется.")
        audio_path = row["audio_path"]
        if not audio_path or not os.path.isfile(audio_path) or os.path.getsize(audio_path) == 0:
            raise HTTPException(400, "Аудиозапись не сохранена или пуста.")
        recording_streams.prepare_completion(conn, meeting_id, request.session_id,
                                             request.chunk_count, request.duration_seconds)
        conn.commit()
    status = start_saved_meeting_generation(meeting_id)
    return {"meeting_id": meeting_id, "status": status}


async def run_recording_pipeline(meeting_id: str):
    await run_generation(meeting_id)


@app.post("/api/v1/meetings/{meeting_id}/transcribe")
async def transcribe_meeting(meeting_id: str, background_tasks: BackgroundTasks, auto_generate: bool = False):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM meetings WHERE id = ?", (meeting_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Meeting not found")
        # Mini retries acknowledge an accepted job without another STT/LLM pass.
        if auto_generate and (row["status"] == "transcribing" or (
            row["status"] in ("generating", "cancelling", "cancelled", "completed", "error") and row["transcript_text"]
        )):
            return {"meeting_id": meeting_id, "status": row["status"]}
        if generation_jobs.active(meeting_id) or row["status"] == "transcribing":
            raise HTTPException(409, "Обработка совещания уже выполняется.")
        audio_path = row["audio_path"]
        if not audio_path or not os.path.exists(audio_path):
            raise HTTPException(status_code=400, detail="No audio file uploaded for this meeting")

        conn.execute("UPDATE meetings SET status = 'transcribing', error_message = '' WHERE id = ?", (meeting_id,))
        conn.commit()

    if auto_generate:
        background_tasks.add_task(run_transcription, meeting_id, audio_path, auto_generate=True)
    else:
        background_tasks.add_task(run_transcription, meeting_id, audio_path)
    return {"meeting_id": meeting_id, "status": "transcribing"}

async def run_transcription(meeting_id: str, audio_path: str, auto_generate: bool = False):
    logger.info("Starting transcription for meeting %s with audio %s", meeting_id, audio_path)
    whisper_mode = get_setting("whisper_mode", "local")

    try:
        with get_db() as conn:
            m_row = conn.execute("SELECT source_language FROM meetings WHERE id = ?", (meeting_id,)).fetchone()
            source_lang = m_row["source_language"] if m_row else "multi"

        if whisper_mode == "local":
            whisper_local_model = get_setting("whisper_local_model", "small")
            whisper_device = get_setting("whisper_device", "auto")
            logger.info("Using local Whisper model '%s' (device: %s)...", whisper_local_model, whisper_device)

            transcribe_res = await asyncio.to_thread(
                whisper_service.transcribe_local,
                audio_path=audio_path,
                model_name=whisper_local_model,
                device=whisper_device,
                language=source_lang,
            )

            formatted_transcript = transcribe_res["text"]
            segments = transcribe_res["segments"]
            total_duration = transcribe_res["duration"]

        else:
            whisper_url = get_setting("whisper_base_url", "http://localhost:8000/v1").rstrip("/")
            whisper_model = get_setting("whisper_model", "whisper-1")
            whisper_key = get_setting("whisper_api_key", "").strip()
            auth_headers = get_auth_headers(whisper_key)

            url = f"{whisper_url}/audio/transcriptions"
            async with httpx.AsyncClient(timeout=600.0) as client:
                with open(audio_path, "rb") as f:
                    files = {"file": (os.path.basename(audio_path), f, "audio/mpeg")}
                    data = {
                        "model": whisper_model,
                        "response_format": "verbose_json",
                    }
                    resp = await client.post(url, headers=auth_headers, files=files, data=data)

                if resp.status_code != 200:
                    with open(audio_path, "rb") as f:
                        files = {"file": (os.path.basename(audio_path), f, "audio/mpeg")}
                        data = {"model": whisper_model}
                        resp = await client.post(url, headers=auth_headers, files=files, data=data)

                if resp.status_code != 200:
                    raise Exception(f"Whisper server returned HTTP {resp.status_code}: {resp.text[:300]}")

                result = resp.json()

            segments = []
            raw_text = result.get("text", "")

            if "segments" in result and isinstance(result["segments"], list):
                for idx, s in enumerate(result["segments"]):
                    start = float(s.get("start", 0) or 0)
                    end = float(s.get("end", 0) or 0)
                    start_str = time.strftime("%H:%M:%S", time.gmtime(max(0.0, start)))
                    speaker = s.get("speaker") or f"Спикер {(idx % 3) + 1}"
                    text = s.get("text", "").strip()
                    segments.append({
                        "index": idx,
                        "speaker": speaker,
                        "text": text,
                        "timestamp_start": round(start, 2),
                        "timestamp_end": round(end, 2),
                        "timestamp_str": f"[{start_str}]",
                    })
            else:
                lines = [l.strip() for l in raw_text.split("\n") if l.strip()]
                if not lines:
                    lines = [raw_text]
                for idx, line in enumerate(lines):
                    segments.append({
                        "index": idx,
                        "speaker": f"Спикер {(idx % 2) + 1}",
                        "text": line,
                        "timestamp_start": round(idx * 10.0, 2),
                        "timestamp_end": round((idx + 1) * 10.0, 2),
                        "timestamp_str": f"[{time.strftime('%H:%M:%S', time.gmtime(idx * 10))}]",
                    })

            formatted_transcript = "\n".join(
                f"[{s['speaker']}] {s['timestamp_str']}\n{s['text']}\n" for s in segments
            ) or raw_text

            total_duration = 0.0
            if result.get("duration"):
                try:
                    total_duration = float(result["duration"])
                except (ValueError, TypeError):
                    pass
            if not total_duration and segments:
                total_duration = max((float(s.get("timestamp_end", 0) or 0) for s in segments), default=0.0)

        with get_db() as conn:
            conn.execute("""
                UPDATE meetings 
                SET status = 'transcribed', transcript_text = ?, transcript_segments = ?, duration_seconds = ?, error_message = ''
                WHERE id = ?
            """, (formatted_transcript, json.dumps(segments, ensure_ascii=False), total_duration, meeting_id))
            conn.commit()

        logger.info("Transcription completed for meeting %s (%d segments, duration: %.1fs)", meeting_id, len(segments), total_duration)

        if auto_generate:
            # Handoff is synchronous: once generating is visible, Stop has an owner.
            start_saved_meeting_generation(meeting_id)

    except Exception as e:
        logger.exception("Transcription failed for meeting %s: %s", meeting_id, e)
        with get_db() as conn:
            conn.execute("""
                UPDATE meetings SET status = 'error', error_message = ? WHERE id = ?
            """, (str(e), meeting_id))
            conn.commit()

@app.post("/api/v1/meetings/{meeting_id}/generate")
async def generate_meeting(meeting_id: str, request: GenerateMeetingRequest | None = None):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM meetings WHERE id = ?", (meeting_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Meeting not found")
        if generation_jobs.active(meeting_id) or row["status"] == "transcribing":
            raise HTTPException(409, "Обработка совещания уже выполняется.")
        if not (row["transcript_text"] or "").strip():
            raise HTTPException(400, "Meeting has no transcript to generate from")
    try:
        route = resolve_generation_route(request)
    except ValueError as error:
        raise HTTPException(422, str(error)) from None
    attempt = start_meeting_generation(meeting_id, route)
    return {"meeting_id": meeting_id, "status": "generating", "generation_attempt": attempt}


@app.get("/api/v1/meetings/{meeting_id}/generation-options")
async def generation_options(meeting_id: str):
    with get_db() as conn:
        row = conn.execute("SELECT generation_attempt FROM meetings WHERE id=?", (meeting_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Meeting not found")
    config = generation_config()
    providers = [{"mode": "api", "provider": "api", "label": "API из настроек", "model": config["model"],
                  "requires_model": True, "supports_models": True}]
    for provider, label in (("opencode", "OpenCode"), ("codex", "Codex CLI"), ("claude", "Claude Code"), ("antigravity", "Antigravity")):
        options = config["cli"].get(provider, {})
        providers.append({"mode": "cli", "provider": provider, "label": label,
                          "model": options.get("model", "") if isinstance(options, dict) else "",
                          "requires_model": provider == "opencode", "supports_models": provider in ("opencode", "antigravity")})
    default = next((choice for choice in providers if choice["mode"] == config["mode"] and choice["provider"] == (config["provider"] if config["mode"] == "cli" else "api")), providers[0])
    return {"default": {key: default[key] for key in ("mode", "provider", "model")},
            "providers": providers, "last_attempt": json.loads(row["generation_attempt"] or "{}") or None}


@app.get("/api/v1/meetings/{meeting_id}/generation-models")
async def generation_models(meeting_id: str, mode: Literal["api", "cli"],
                            provider: Optional[Literal["api", "opencode", "codex", "claude", "antigravity"]] = None):
    with get_db() as conn:
        if not conn.execute("SELECT id FROM meetings WHERE id=?", (meeting_id,)).fetchone():
            raise HTTPException(404, "Meeting not found")
    config = generation_config()
    if mode == "api":
        if provider not in (None, "api"):
            raise HTTPException(422, "Для режима API выберите настроенный API-провайдер")
        models = await discover_api_models(config["base_url"].rstrip("/"), config["api_key"])
        return {"provider": "api", "models": [{"id": model, "name": model, "provider": "api", "free": None} for model in models],
                "pricing_available": False, "fetched_at": datetime.now(timezone.utc).isoformat()}
    provider = provider or config["provider"]
    if provider not in cli_llm.PROVIDERS:
        raise HTTPException(422, "Неизвестный CLI провайдер")
    options = config["cli"].get(provider, {})
    return await cli_llm.models(provider, (options.get("path", "") or "") if isinstance(options, dict) else "")


@app.post("/api/v1/meetings/{meeting_id}/stop-generation")
async def stop_generation(meeting_id: str):
    with get_db() as conn:
        if not conn.execute("SELECT id FROM meetings WHERE id=?", (meeting_id,)).fetchone():
            raise HTTPException(404, "Meeting not found")
    try:
        await generation_jobs.stop(meeting_id)
    except TimeoutError as error:
        raise HTTPException(504, str(error)) from error
    except RuntimeError as error:
        raise HTTPException(500, str(error)) from error
    with get_db() as conn:
        status = conn.execute("SELECT status FROM meetings WHERE id=?", (meeting_id,)).fetchone()["status"]
    return {"meeting_id": meeting_id, "status": status}


def save_generation_results(meeting_id, results):
    """Publish one coherent generation; cancellation/errors retain the prior set."""
    generation_jobs.check(meeting_id)
    fields = ("protocol_ru", "protocol_kz", "summary_ru", "summary_kz", "template_values_ru", "template_values_kz")
    if set(results) != set(fields):
        raise ValueError("Incomplete generation results")
    with get_db() as conn:
        conn.execute("""UPDATE meetings SET protocol_ru=?, protocol_kz=?, summary_ru=?, summary_kz=?,
                     template_values_ru=?, template_values_kz=?, status='completed', error_message='' WHERE id=?""",
                     (*[json.dumps(results[field], ensure_ascii=False) for field in fields], meeting_id))
        conn.commit()
    generation_jobs.mark_committed(meeting_id)


async def run_generation(meeting_id: str):
    generation_jobs.check(meeting_id)
    logger.info("Starting protocol & summary generation for meeting %s", meeting_id)
    with get_db() as conn:
        row = conn.execute("SELECT * FROM meetings WHERE id = ?", (meeting_id,)).fetchone()
        if not row:
            return

    transcript_text = row["transcript_text"]
    agenda = row["agenda"] or None
    raw_participants = json.loads(row["participants"] or "[]")
    participants = [
        f"{p.get('name', '')} ({p.get('position', '')})" if p.get('position') else p.get('name', '')
        for p in raw_participants if p.get('name')
    ] or None

    generation_jobs.check(meeting_id)
    # Step 1: Russian Protocol
    logger.info("Generating Russian protocol...")
    protocol_ru = await protocol.generate(
        transcript_text=transcript_text,
        participants=participants,
        agenda=agenda,
    )


    generation_jobs.check(meeting_id)
    # Step 2: Kazakh Protocol Translation
    logger.info("Translating protocol to Kazakh...")
    protocol_kz = await protocol._translate(
        protocol_ru=protocol_ru,
        system_prompt=prompts.TRANSLATE_PROTOCOL_SYSTEM,
        language_name="Kazakh",
        has_agenda=agenda is not None,
    )


    generation_jobs.check(meeting_id)
    # Step 3: Russian Summary
    logger.info("Generating Russian summary...")
    summary_ru = await summary.generate(transcript_text=transcript_text)


    generation_jobs.check(meeting_id)
    # Step 4: Kazakh Summary Translation
    logger.info("Translating summary to Kazakh...")
    summary_kz = await summary.translate(summary_ru=summary_ru)


    tpl_values_ru, tpl_values_kz = {}, {}
    # Generate template values if meeting is attached to a template
    tpl_id = row["template_id"] if "template_id" in row.keys() else ""
    if tpl_id:
        try:
            generation_jobs.check(meeting_id)
            with get_db() as conn:
                tpl_row = conn.execute("SELECT * FROM protocol_templates WHERE id = ?", (tpl_id,)).fetchone()
            if tpl_row and os.path.exists(tpl_row["docx_path"]):
                with open(tpl_row["docx_path"], "rb") as f:
                    tpl_bytes = f.read()
                parsed_tpl = docx_template.parse_docx_template(tpl_bytes)
                slots_data = json.loads(tpl_row["slots"] or "[]")
                stored_slots_map = {s["key"]: s for s in slots_data}
                loop_iters = docx_template._loop_iterables(tpl_bytes)
                for s in parsed_tpl.slots:
                    st = stored_slots_map.get(s.key)
                    if st:
                        s.label = st.get("label", s.label)
                        s.value_type = st.get("value_type", s.value_type)
                        s.repeat = st.get("repeat", s.repeat)
                        s.omit_when_empty = st.get("omit_when_empty", s.omit_when_empty)
                    elif s.key in loop_iters:
                        s.value_type = "list[object]"

                logger.info("Generating custom template values for meeting %s with template %s", meeting_id, tpl_id)
                protocol_source = template_protocol.build_protocol_source(
                    protocol_ru, agenda=agenda, participants=participants,
                )
                transfer_res = await template_generation.generate_protocol_template_values(
                    parsed_tpl,
                    tpl_row["additional_prompt"] or "",
                    transcript_text,
                    protocol_source=protocol_source,
                    output_language="Russian",
                    detail_level=tpl_row["detail_level"] or "concise",
                )
                tpl_values_ru = transfer_res.get("template_values") or {}

                try:
                    tpl_values_kz = await template_generation.translate_template_values(
                        parsed_tpl, tpl_values_ru, "Kazakh",
                    )
                except Exception as kz_err:
                    logger.warning("Template values KZ translation error: %s", kz_err)
                    raise

        except Exception as tpl_err:
            logger.warning("Custom template generation failed: %s", tpl_err)
            raise

    save_generation_results(meeting_id, {
        "protocol_ru": protocol_ru, "protocol_kz": protocol_kz,
        "summary_ru": summary_ru, "summary_kz": summary_kz,
        "template_values_ru": tpl_values_ru, "template_values_kz": tpl_values_kz,
    })
    logger.info("Meeting generation completed successfully for %s", meeting_id)

    # Auto-index into local vector store for RAG
    try:
        generation_jobs.check(meeting_id)
        with get_db() as conn:
            await indexer.index_meeting(meeting_id, conn, check_current=lambda: generation_jobs.check(meeting_id))
        logger.info("Auto-indexed meeting %s into vector store", meeting_id)
    except Exception as idx_err:
        logger.warning("Failed to auto-index meeting %s: %s", meeting_id, idx_err)

@app.get("/api/v1/meetings")
async def list_meetings(q: Optional[str] = None):
    with get_db() as conn:
        if q:
            query = "%" + q + "%"
            rows = conn.execute("""
                SELECT id, title, created_at, status, duration_seconds, source_language, error_message
                FROM meetings
                WHERE title LIKE ? OR agenda LIKE ? OR transcript_text LIKE ?
                ORDER BY created_at DESC
            """, (query, query, query)).fetchall()
        else:
            rows = conn.execute("""
                SELECT id, title, created_at, status, duration_seconds, source_language, error_message
                FROM meetings
                ORDER BY created_at DESC
            """).fetchall()

        return [dict(r) for r in rows]

@app.get("/api/v1/meetings/{meeting_id}")
async def get_meeting(meeting_id: str):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM meetings WHERE id = ?", (meeting_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Meeting not found")

        data = dict(row)
        data["participants"] = json.loads(data["participants"] or "[]")
        data["transcript_segments"] = json.loads(data["transcript_segments"] or "[]")
        data["protocol_ru"] = json.loads(data["protocol_ru"] or "{}")
        data["protocol_kz"] = json.loads(data["protocol_kz"] or "{}")
        data["summary_ru"] = json.loads(data["summary_ru"] or "{}")
        data["summary_kz"] = json.loads(data["summary_kz"] or "{}")
        data["template_values_ru"] = json.loads(data.get("template_values_ru") or "{}")
        data["template_values_kz"] = json.loads(data.get("template_values_kz") or "{}")
        data["template_id"] = data.get("template_id") or ""
        data["generation_attempt"] = json.loads(data.get("generation_attempt") or "{}") or None
        return data

@app.patch("/api/v1/meetings/{meeting_id}")
async def update_meeting(meeting_id: str, req: UpdateMeetingRequest):
    with get_db() as conn:
        row = conn.execute("SELECT id FROM meetings WHERE id = ?", (meeting_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Meeting not found")

        updates = []
        params = []
        if req.title is not None:
            updates.append("title = ?")
            params.append(req.title)
        if req.source_language is not None:
            updates.append("source_language = ?")
            params.append(req.source_language)
        if req.agenda is not None:
            updates.append("agenda = ?")
            params.append(req.agenda)
        if req.participants is not None:
            updates.append("participants = ?")
            params.append(json.dumps(req.participants, ensure_ascii=False))
        if req.transcript_text is not None:
            updates.append("transcript_text = ?")
            params.append(req.transcript_text)
        if req.protocol_ru is not None:
            updates.append("protocol_ru = ?")
            params.append(json.dumps(req.protocol_ru, ensure_ascii=False))
        if req.protocol_kz is not None:
            updates.append("protocol_kz = ?")
            params.append(json.dumps(req.protocol_kz, ensure_ascii=False))
        if req.summary_ru is not None:
            updates.append("summary_ru = ?")
            params.append(json.dumps(req.summary_ru, ensure_ascii=False))
        if req.summary_kz is not None:
            updates.append("summary_kz = ?")
            params.append(json.dumps(req.summary_kz, ensure_ascii=False))
        if req.template_id is not None:
            updates.append("template_id = ?")
            params.append(req.template_id)

        if updates:
            params.append(meeting_id)
            conn.execute(f"UPDATE meetings SET {', '.join(updates)} WHERE id = ?", params)
            conn.commit()

        return {"success": True}

@app.delete("/api/v1/meetings/{meeting_id}")
async def delete_meeting(meeting_id: str):
    with get_db() as conn:
        row = conn.execute("SELECT audio_path FROM meetings WHERE id = ?", (meeting_id,)).fetchone()
        if row and row["audio_path"] and os.path.exists(row["audio_path"]):
            try:
                os.remove(row["audio_path"])
            except Exception:
                pass
        recording_streams.discard_meeting(conn, meeting_id)
        conn.execute("DELETE FROM meetings WHERE id = ?", (meeting_id,))
        conn.commit()
    return {"success": True}


# ---------------------------------------------------------------------------
# Template Helper Utilities
# ---------------------------------------------------------------------------

def _mock_template_values(slots: list[Any]) -> dict[str, Any]:
    """Fallback values generator for previewing/testing templates even when LLM is offline."""
    values: dict[str, Any] = {}
    now = datetime.now()
    for s in slots:
        key = s.key if hasattr(s, "key") else s.get("key", "")
        val_type = s.value_type if hasattr(s, "value_type") else s.get("value_type", "string")
        repeat = s.repeat if hasattr(s, "repeat") else s.get("repeat")
        label = s.label if hasattr(s, "label") else s.get("label", key)

        if val_type == "date":
            values[key] = now.strftime("%d.%m.%Y")
        elif repeat and repeat.get("kind") == "numbered_outline":
            values[key] = [
                {"level": "0", "text": "Обсуждение текущих вопросов повестки"},
                {"level": "0", "text": "Утверждение проектных решений и задач"},
                {"level": "0", "text": "Назначение контрольных сроков"},
            ]
        elif repeat and repeat.get("kind") == "table_rows":
            cols = [c["key"] for c in repeat.get("columns", [])]
            row1, row2, row3 = {}, {}, {}
            for c in cols:
                row1[c] = "Ахметов Б.С." if "fio" in c else ("Председатель" if "dolzh" in c or "rol" in c else f"Значение {c}")
                row2[c] = "Калиева Д.С." if "fio" in c else ("Секретарь" if "dolzh" in c or "rol" in c else f"Значение {c}")
                row3[c] = "Сергеев В.П." if "fio" in c else ("Член комиссии" if "dolzh" in c or "rol" in c else f"Значение {c}")
            values[key] = [row1, row2, row3]
        elif val_type == "list[object]":
            values[key] = [
                {"фио": "Ахметов Б.С.", "должность": "Председатель правления"},
                {"фио": "Калиева Д.С.", "должность": "Секретарь"},
                {"фио": "Сергеев В.П.", "должность": "Технический директор"},
            ]
        elif val_type == "list[string]":
            values[key] = [
                "1. Утвердить отчет и принять план работ.",
                "2. Завершить тестирование системы до 15 числа.",
            ]
        else:
            k = key.lower()
            if "председатель_фио" in k or "predsedatel" in k:
                values[key] = "Ахметов Б.С."
            elif "секретарь_фио" in k or "sekretar" in k:
                values[key] = "Калиева Д.С."
            elif "председатель_должность" in k:
                values[key] = "Председатель правления"
            elif "город" in k or "gorod" in k:
                values[key] = "Астана"
            elif "день" in k or "den" in k:
                values[key] = now.strftime("%d")
            elif "месяц" in k or "mesyac" in k:
                months = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа", "сентября", "октября", "ноября", "декабря"]
                values[key] = months[now.month - 1]
            elif "год" in k or "god" in k:
                values[key] = now.strftime("%y")
            elif "время" in k or "vremya" in k:
                values[key] = now.strftime("%H:%M")
            elif "номер" in k:
                values[key] = "14-ПР"
            elif "следующ" in k:
                values[key] = (now + timedelta(days=30)).strftime("%d.%m.%Y")
            elif "reshili" in k or "постанов" in k or "реш" in k:
                values[key] = "1. Одобрить проект решения единогласно. 2. Утвердить план реализации."
            elif "golosov_za" in k:
                values[key] = "5"
            elif "golosov_protiv" in k or "golosov_vozderzhalis" in k:
                values[key] = "0"
            elif "mesto" in k or "место" in k:
                values[key] = "г. Астана, головной офис"
            elif "povestka" in k or "вопрос" in k:
                values[key] = "Обсуждение ключевых задач и утверждение плана мероприятий"
            else:
                values[key] = f"Заполнено ({label})"
    return values


def _adapt_meeting_to_template_values(slots: list[Any], meeting_row: dict[str, Any], protocol_data: dict[str, Any], lang: str = "ru") -> dict[str, Any]:
    """Fill template slots from meeting protocol data."""
    stored_values = meeting_row.get("template_values_kz" if lang == "kz" else "template_values_ru")
    if stored_values and isinstance(stored_values, str):
        try:
            stored_values = json.loads(stored_values)
        except Exception:
            stored_values = {}
    if stored_values and isinstance(stored_values, dict) and len(stored_values) > 0:
        return stored_values

    # Base values from fallback generator
    values = _mock_template_values(slots)

    # Enrich from actual meeting metadata and protocol
    title = meeting_row.get("title", "")
    created_at_str = meeting_row.get("created_at", "")[:10]
    agenda = meeting_row.get("agenda", "") or protocol_data.get("metadata", {}).get("agenda", "")

    try:
        meeting_participants = json.loads(meeting_row.get("participants") or "[]")
    except Exception:
        meeting_participants = []

    topics = protocol_data.get("agenda_items") or protocol_data.get("topics") or []
    all_decisions = []
    for t in topics:
        for d in t.get("decisions", []):
            if isinstance(d, dict):
                text = d.get("decision", "")
                resp = d.get("responsible", "")
                dl = d.get("deadline", "")
                meta = []
                if resp: meta.append(f"Отв: {resp}")
                if dl: meta.append(f"Срок: {dl}")
                all_decisions.append(f"{text} ({', '.join(meta)})" if meta else text)
            elif isinstance(d, str):
                all_decisions.append(d)

    for s in slots:
        key = s.key if hasattr(s, "key") else s.get("key", "")
        repeat = s.repeat if hasattr(s, "repeat") else s.get("repeat")
        k = key.lower()

        if repeat and repeat.get("kind") == "table_rows":
            cols = [c["key"] for c in repeat.get("columns", [])]
            if meeting_participants:
                rows = []
                for p in meeting_participants:
                    name = p.get("name", "") if isinstance(p, dict) else str(p)
                    pos = p.get("position", "") if isinstance(p, dict) else ""
                    role = p.get("role", "Член комиссии") if isinstance(p, dict) else "Участник"
                    row_obj = {}
                    for c in cols:
                        if "fio" in c: row_obj[c] = name
                        elif "dolzh" in c: row_obj[c] = pos or "Сотрудник"
                        elif "rol" in c: row_obj[c] = role
                        else: row_obj[c] = ""
                    rows.append(row_obj)
                values[key] = rows

        elif repeat and repeat.get("kind") == "numbered_outline":
            if topics:
                outline_rows = []
                for t in topics:
                    t_name = t.get("topic") or t.get("topic_name") or ""
                    outline_rows.append({"level": "0", "text": t_name})
                    for dec in t.get("decisions", []):
                        d_text = dec.get("decision", "") if isinstance(dec, dict) else str(dec)
                        if d_text:
                            outline_rows.append({"level": "1", "text": f"Решение: {d_text}"})
                if outline_rows:
                    values[key] = outline_rows

        elif k == "участники" or (isinstance(values.get(key), list) and len(values.get(key, [])) > 0 and isinstance(values[key][0], dict) and "фио" in values[key][0]):
            if meeting_participants:
                obj_list = []
                for p in meeting_participants:
                    p_name = p.get("name", "") if isinstance(p, dict) else str(p)
                    p_pos = p.get("position", "") if isinstance(p, dict) else "Участник"
                    obj_list.append({"фио": p_name, "должность": p_pos})
                values[key] = obj_list

        elif "reshili" in k or "постанов" in k or "реш" in k:
            if all_decisions:
                values[key] = "\n".join(f"{i+1}. {d}" for i, d in enumerate(all_decisions))
        elif "povestka" in k or "вопрос" in k:
            if agenda:
                values[key] = agenda
        elif "title" in k or "наименование" in k:
            if title:
                values[key] = title

    return values


# ---------------------------------------------------------------------------
# Protocol Templates Endpoints
# ---------------------------------------------------------------------------

@app.get("/api/v1/templates")
async def list_templates():
    with get_db() as conn:
        rows = conn.execute("""
            SELECT id, name, description, additional_prompt, detail_level, status,
                   slots, render_ready, test_docx_path, created_at
            FROM protocol_templates
            ORDER BY created_at DESC
        """).fetchall()

        result = []
        for r in rows:
            data = dict(r)
            try:
                data["slots"] = json.loads(data["slots"] or "[]")
            except Exception:
                data["slots"] = []
            data["slots_count"] = len(data["slots"])
            data["has_test_docx"] = bool(data.get("test_docx_path") and os.path.exists(data["test_docx_path"]))
            result.append(data)
        return result


@app.get("/api/v1/templates/sample/download")
async def download_sample_template():
    candidate_paths = [
        BASE_DIR.parent / "public" / "templates" / "ideal-protocol-template.docx",
        BASE_DIR.parent.parent / "jynalys" / "jinalys-new-admin-frontend" / "public" / "templates" / "ideal-protocol-template.docx",
    ]
    sample_path = None
    for p in candidate_paths:
        if p.exists():
            sample_path = p
            break
    if not sample_path:
        raise HTTPException(status_code=404, detail="Sample template docx not found")
    return FileResponse(
        str(sample_path),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="ideal-protocol-template.docx",
    )


@app.get("/api/v1/templates/{template_id}")
async def get_template_detail(template_id: str):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM protocol_templates WHERE id = ?", (template_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Template not found")

        data = dict(row)
        data["slots"] = json.loads(data["slots"] or "[]")
        data["schema_json"] = json.loads(data["schema_json"] or "{}")
        data["style_config"] = json.loads(data["style_config"] or "{}")
        data["template_profile"] = json.loads(data["template_profile"] or "{}")
        data["test_values"] = json.loads(data.get("test_values") or "{}")
        data["has_test_docx"] = bool(data.get("test_docx_path") and os.path.exists(data["test_docx_path"]))
        data["has_test_pdf"] = bool(data.get("test_pdf_path") and os.path.exists(data["test_pdf_path"]))
        return data


@app.post("/api/v1/templates/parse-preview")
async def preview_template_slots(file: UploadFile = File(...)):
    """Validate and parse an uploaded docx without saving it."""
    content = await file.read()
    if len(content) > 50 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Размер файла превышает лимит 50 МБ")

    try:
        working_bytes, parsed = docx_template.prepare_template_pipeline(content)
        name = Path(file.filename or "template").stem
        descriptor = docx_template.parsed_to_descriptor(parsed, name)
        return {
            "slots": descriptor["slots"],
            "render_ready": descriptor["render_ready"],
            "stats": descriptor["stats"],
            "warnings": descriptor["warnings"],
        }
    except Exception as e:
        logger.exception("Parse preview failed: %s", e)
        raise HTTPException(status_code=400, detail=f"Ошибка разбора DOCX: {e}")


@app.post("/api/v1/templates/upload")
async def upload_template(
    file: UploadFile = File(...),
    name: str = Form(...),
    description: Optional[str] = Form(""),
    additional_prompt: Optional[str] = Form(""),
    detail_level: Optional[str] = Form("concise"),
):
    sync_generation_settings()
    docx_bytes = await file.read()
    if len(docx_bytes) > 50 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Размер файла превышает лимит 50 МБ")

    try:
        working_bytes, parsed = docx_template.prepare_template_pipeline(docx_bytes)
        if not parsed.slots:
            raise HTTPException(
                status_code=400,
                detail="В документе не найдены поля для заполнения. Используйте {{ field_name }}, линии ___ или инструкции [в скобках]."
            )
        descriptor = docx_template.parsed_to_descriptor(parsed, name)
        profile_data = await template_profile.build_template_profile(working_bytes, parsed, language="ru")
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("DOCX parsing error: %s", e)
        raise HTTPException(status_code=400, detail=f"Ошибка обработки DOCX: {e}")

    template_id = str(uuid.uuid4())
    working_path = TEMPLATES_DIR / f"{template_id}.docx"
    source_path = TEMPLATES_DIR / f"{template_id}_source.docx"

    with open(working_path, "wb") as f:
        f.write(working_bytes)
    with open(source_path, "wb") as f:
        f.write(docx_bytes)

    now = datetime.now(timezone.utc).isoformat()
    with get_db() as conn:
        conn.execute("""
            INSERT INTO protocol_templates (
                id, name, description, additional_prompt, detail_level, status,
                docx_path, source_docx_path, slots, schema_json, style_config,
                template_profile, render_ready, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            template_id,
            name.strip() or "Пользовательский шаблон",
            (description or "").strip(),
            (additional_prompt or "").strip(),
            (detail_level or "concise").strip(),
            "approved",
            str(working_path),
            str(source_path),
            json.dumps(descriptor["slots"], ensure_ascii=False),
            json.dumps(descriptor["schema_json"], ensure_ascii=False),
            json.dumps(descriptor["style_config"], ensure_ascii=False),
            json.dumps(profile_data, ensure_ascii=False),
            1 if descriptor["render_ready"] else 0,
            now,
        ))
        conn.commit()

    return {
        "id": template_id,
        "name": name,
        "slots_count": len(descriptor["slots"]),
        "render_ready": descriptor["render_ready"],
        "created_at": now,
    }


@app.delete("/api/v1/templates/{template_id}")
async def delete_template(template_id: str):
    if template_id == "default-protocol-template":
        raise HTTPException(status_code=403, detail="Системный шаблон удалять нельзя")
    with get_db() as conn:
        row = conn.execute("SELECT docx_path, source_docx_path, test_docx_path, test_pdf_path FROM protocol_templates WHERE id = ?", (template_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Template not found")

        # Detach meetings bound to this template so export falls back to standard protocol
        detached = 0
        try:
            cnt_row = conn.execute("SELECT COUNT(*) AS cnt FROM meetings WHERE template_id = ?", (template_id,)).fetchone()
            detached = int(cnt_row["cnt"]) if cnt_row else 0
            if detached:
                conn.execute("UPDATE meetings SET template_id = '' WHERE template_id = ?", (template_id,))
        except Exception:
            logger.warning("Could not detach meetings from template %s", template_id)
            detached = 0

        for p in [row["docx_path"], row["source_docx_path"], row["test_docx_path"], row["test_pdf_path"]]:
            if p and os.path.exists(p):
                try: os.remove(p)
                except Exception: pass

        conn.execute("DELETE FROM protocol_templates WHERE id = ?", (template_id,))
        conn.commit()
    return {"success": True, "detached_meetings": detached}


@app.get("/api/v1/templates/{template_id}/download")
async def download_template_file(template_id: str, variant: str = Query("working", pattern="^(working|source)$")):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM protocol_templates WHERE id = ?", (template_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Template not found")

    file_path = row["source_docx_path"] if variant == "source" else row["docx_path"]
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="File on disk not found")

    safe_name = f"{row['name']}_{variant}.docx"
    return FileResponse(
        file_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=safe_name,
    )


@app.post("/api/v1/templates/{template_id}/test")
async def test_template_run(template_id: str, req: Optional[TemplateTestRequest] = None):
    sync_generation_settings()
    with get_db() as conn:
        row = conn.execute("SELECT * FROM protocol_templates WHERE id = ?", (template_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Template not found")

    if not os.path.exists(row["docx_path"]):
        raise HTTPException(status_code=404, detail="Template file not found on disk")

    with open(row["docx_path"], "rb") as f:
        docx_bytes = f.read()

    slots_data = json.loads(row["slots"] or "[]")
    working_bytes, parsed = docx_template.prepare_template_pipeline(docx_bytes, stored_slots=slots_data)

    test_transcript = (req.transcript.strip() if req and req.transcript else "") or template_generation.DEFAULT_TEST_TRANSCRIPT
    dl = (req.detail_level if req and req.detail_level else "") or row["detail_level"] or "concise"

    tpl_profile = json.loads(row["template_profile"] or "{}")
    if not tpl_profile:
        tpl_profile = template_profile.heuristic_template_profile(parsed, language="ru")

    generation_started = time.perf_counter()
    values = None
    try:
        gen_res = await template_protocol.generate(
            parsed=parsed,
            transcript=test_transcript,
            additional_prompt=row["additional_prompt"] or "",
            detail_level=dl,
            template_profile=tpl_profile,
            output_language="Russian",
        )
        values = gen_res.get("template_values") or {}
    except Exception as e:
        logger.warning("template_protocol.generate failed (%s), falling back to template_generation.generate_template_values", e)
        try:
            values = await template_generation.generate_template_values(
                parsed=parsed,
                additional_prompt=row["additional_prompt"] or "",
                transcript=test_transcript,
                output_language="Russian",
                detail_level=dl,
            )
        except Exception as gen_err:
            logger.exception("LLM template test failed for template %s", template_id)
            raise HTTPException(status_code=502, detail=f"Модель не смогла сформировать тестовый протокол: {gen_err}") from gen_err

    try:
        rendered = docx_template.render_docx_template(working_bytes, values, slots=parsed.slots)
    except Exception as e:
        logger.exception("Render test docx failed for template %s", template_id)
        raise HTTPException(status_code=422, detail=f"Модель вернула данные, которые не удалось вставить в DOCX: {e}") from e

    test_docx_path = TEMPLATES_DIR / f"{template_id}_test.docx"
    with open(test_docx_path, "wb") as f:
        f.write(rendered)

    test_pdf_path = TEMPLATES_DIR / f"{template_id}_test.pdf"
    try:
        descriptor = docx_template.parsed_to_descriptor(parsed, row["name"])
        test_pdf_bytes = pdf_export.build_template_test_pdf(row["name"], values, descriptor["slots"])
        with open(test_pdf_path, "wb") as f:
            f.write(test_pdf_bytes)
    except Exception as pdf_err:
        logger.warning("Failed to generate test PDF for template %s: %s", template_id, pdf_err)

    has_pdf = os.path.exists(test_pdf_path)
    with get_db() as conn:
        conn.execute("""
            UPDATE protocol_templates
            SET test_values = ?, test_docx_path = ?, test_pdf_path = ?
            WHERE id = ?
        """, (
            json.dumps(values, ensure_ascii=False),
            str(test_docx_path),
            str(test_pdf_path) if has_pdf else "",
            template_id,
        ))
        conn.commit()

    return {
        "template_id": template_id,
        "values": values,
        "slots_filled": len(values),
        "has_test_docx": True,
        "has_test_pdf": has_pdf,
        "model": (get_cli_config().get(settings.LLM_CLI_PROVIDER, {}).get("model") or "CLI default") if settings.LLM_MODE == "cli" else get_setting("llm_model", settings.LLM_MODEL),
        "generation_seconds": round(time.perf_counter() - generation_started, 2),
    }


@app.get("/api/v1/templates/{template_id}/test-download")
async def download_template_test_file(template_id: str):
    with get_db() as conn:
        row = conn.execute("SELECT name, test_docx_path FROM protocol_templates WHERE id = ?", (template_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Template not found")

    if not row["test_docx_path"] or not os.path.exists(row["test_docx_path"]):
        raise HTTPException(status_code=404, detail="Тестовый DOCX еще не сгенерирован. Запустите тестирование.")

    return FileResponse(
        row["test_docx_path"],
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=f"test_{row['name']}.docx",
    )


@app.get("/api/v1/templates/{template_id}/test-download-pdf")
async def download_template_test_pdf(template_id: str):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM protocol_templates WHERE id = ?", (template_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Template not found")

    pdf_path = row["test_pdf_path"] if "test_pdf_path" in row.keys() else ""
    if pdf_path and os.path.exists(pdf_path):
        return FileResponse(
            pdf_path,
            media_type="application/pdf",
            filename=f"test_{row['name']}.pdf",
        )

    # If test_values exist, render on the fly
    test_vals = json.loads(row["test_values"] or "{}")
    if not test_vals:
        raise HTTPException(status_code=404, detail="Тестовый протокол еще не сгенерирован. Запустите тестирование.")

    slots = json.loads(row["slots"] or "[]")
    pdf_bytes = pdf_export.build_template_test_pdf(row["name"], test_vals, slots)
    clean_name = re.sub(r'[/\\?%*:|"<>]+', '-', row["name"]).strip() or "template"
    utf8_filename = f"test_{clean_name}.pdf"
    quoted = urllib.parse.quote(utf8_filename)
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'attachment; filename="test_{template_id[:8]}.pdf"; filename*=UTF-8\'\'{quoted}'
        },
    )


# ---------------------------------------------------------------------------
# AI Assistant & Chat Endpoints
# ---------------------------------------------------------------------------

@app.post("/api/v1/chat")
async def chat_assistant(req: ChatRequest):
    sync_generation_settings()
    context_parts = [
        "You are an expert AI meeting assistant. Answer questions concisely and professionally based on the meeting data.",
    ]

    if req.meeting_id:
        with get_db() as conn:
            row = conn.execute("SELECT * FROM meetings WHERE id = ?", (req.meeting_id,)).fetchone()
            if row:
                if row["transcript_text"]:
                    context_parts.append(f"\n\nTranscript:\n{row['transcript_text'][:6000]}")
                if row["protocol_ru"] and row["protocol_ru"] != "{}":
                    context_parts.append(f"\n\nProtocol:\n{row['protocol_ru'][:4000]}")
                if row["summary_ru"] and row["summary_ru"] != "{}":
                    context_parts.append(f"\n\nSummary:\n{row['summary_ru'][:2000]}")

    if req.context:
        if req.context.get("protocol"):
            context_parts.append(f"\n\nProtocol Context:\n{req.context['protocol'][:4000]}")
        if req.context.get("summary"):
            context_parts.append(f"\n\nSummary Context:\n{req.context['summary'][:2000]}")

    system_prompt = "\n".join(context_parts)
    messages = [{"role": "system", "content": system_prompt}] + [
        {"role": m.role, "content": m.content} for m in req.messages
    ]

    async def event_generator():
        try:
            async for delta in llm.stream_messages(messages, temperature=0.3):
                frame = {"choices": [{"delta": {"content": delta}}]}
                yield f"data: {json.dumps(frame, ensure_ascii=False)}\n\n"
            yield "data: [DONE]\n\n"
        except Exception as e:
            yield f"data: {json.dumps({'error': str(e)})}\n\n"

    return StreamingResponse(event_generator(), media_type="text/event-stream")

# ---------------------------------------------------------------------------
# Multi-Turn RAG Chat Sessions & Vector Search API
# ---------------------------------------------------------------------------

class CreateChatSessionRequest(BaseModel):
    title: Optional[str] = "Новый диалог"

class SendSessionMessageRequest(BaseModel):
    content: str
    filter_meeting_id: Optional[str] = None

@app.get("/api/v1/chat/sessions")
async def list_chat_sessions():
    with get_db() as conn:
        rows = conn.execute("""
            SELECT s.id, s.title, s.created_at, s.updated_at,
                   COUNT(m.id) as message_count,
                   (SELECT content FROM chat_messages WHERE session_id = s.id ORDER BY created_at DESC LIMIT 1) as last_message
            FROM chat_sessions s
            LEFT JOIN chat_messages m ON m.session_id = s.id
            GROUP BY s.id
            ORDER BY s.updated_at DESC
        """).fetchall()
        return [dict(r) for r in rows]

@app.post("/api/v1/chat/sessions")
async def create_chat_session(req: Optional[CreateChatSessionRequest] = None):
    session_id = str(uuid.uuid4())
    title = (req.title if req and req.title else "Новый диалог").strip()
    with get_db() as conn:
        conn.execute(
            "INSERT INTO chat_sessions (id, title) VALUES (?, ?)",
            (session_id, title)
        )
        conn.commit()
        row = conn.execute("SELECT * FROM chat_sessions WHERE id = ?", (session_id,)).fetchone()
        return dict(row)

@app.get("/api/v1/chat/sessions/{session_id}/messages")
async def get_session_messages(session_id: str):
    with get_db() as conn:
        rows = conn.execute(
            "SELECT * FROM chat_messages WHERE session_id = ? ORDER BY created_at ASC",
            (session_id,)
        ).fetchall()
        result = []
        for r in rows:
            item = dict(r)
            if item.get("citations"):
                try:
                    item["citations"] = json.loads(item["citations"])
                except Exception:
                    item["citations"] = []
            else:
                item["citations"] = []
            result.append(item)
        return result

@app.delete("/api/v1/chat/sessions/{session_id}")
async def delete_chat_session(session_id: str):
    with get_db() as conn:
        conn.execute("DELETE FROM chat_messages WHERE session_id = ?", (session_id,))
        conn.execute("DELETE FROM chat_sessions WHERE id = ?", (session_id,))
        conn.commit()
    return {"success": True}

@app.post("/api/v1/chat/sessions/{session_id}/messages")
async def send_session_message(session_id: str, req: SendSessionMessageRequest):
    content = req.content.strip()
    if not content:
        raise HTTPException(status_code=400, detail="Content cannot be empty")

    with get_db() as conn:
        session = conn.execute("SELECT * FROM chat_sessions WHERE id = ?", (session_id,)).fetchone()
        if not session:
            raise HTTPException(status_code=404, detail="Chat session not found")

        # Stage the user message until generation succeeds. Failed/cancelled
        # attempts can be retried without duplicating persisted conversation.
        user_msg_id = str(uuid.uuid4())

        # 3. Retrieve context & citations via RAG
        top_chunks, citations = await retriever.retrieve(
            conn=conn,
            query=content,
            top_k=6,
            filter_meeting_id=req.filter_meeting_id,
        )

        # 4. Fetch recent history for multi-turn context
        history_rows = conn.execute(
            "SELECT role, content FROM chat_messages WHERE session_id = ? ORDER BY created_at ASC",
            (session_id,)
        ).fetchall()
        history = [dict(r) for r in history_rows]

    # 5. Build prompt with RAG context
    prompt = retriever.build_rag_prompt(history, top_chunks, content)

    # 6. Stream SSE events
    sync_generation_settings()

    cli_mode = settings.LLM_MODE == "cli"

    async def event_generator():
        # Frame 1: Citations event
        yield f"event: citations\ndata: {json.dumps({'citations': citations}, ensure_ascii=False)}\n\n"

        assistant_text = ""
        try:
            async for delta in llm.stream_chat([{"role": "user", "content": prompt}], temperature=0.3):
                if delta:
                    assistant_text += delta
                    yield f"event: delta\ndata: {json.dumps({'delta': delta}, ensure_ascii=False)}\n\n"

            # Save assistant message to database
            asst_msg_id = str(uuid.uuid4())
            with get_db() as conn:
                conn.execute(
                    "INSERT INTO chat_messages (id, session_id, role, content, citations) VALUES (?, ?, ?, ?, ?)",
                    (user_msg_id, session_id, "user", content, None),
                )
                if session["title"] in ("Новый диалог", "New Chat", ""):
                    conn.execute("UPDATE chat_sessions SET title = ? WHERE id = ?",
                                 (content[:45] + ("..." if len(content) > 45 else ""), session_id))
                conn.execute(
                    "INSERT INTO chat_messages (id, session_id, role, content, citations) VALUES (?, ?, ?, ?, ?)",
                    (asst_msg_id, session_id, "assistant", assistant_text, json.dumps(citations, ensure_ascii=False))
                )
                conn.execute("UPDATE chat_sessions SET updated_at = CURRENT_TIMESTAMP WHERE id = ?", (session_id,))
                conn.commit()

            yield f"event: done\ndata: {json.dumps({'message_id': asst_msg_id, 'full_text': assistant_text}, ensure_ascii=False)}\n\n"

        except Exception as e:
            logger.warning("LLM generation stream failed: %s", e)
            if cli_mode:
                yield f"event: error\ndata: {json.dumps({'error': str(e)}, ensure_ascii=False)}\n\n"
                return
            if not assistant_text:
                if citations:
                    fallback_text = (
                        "⚠️ **Локальная языковая модель не запущена или не загружена в Ollama.**\n\n"
                        "Тем не менее, векторный поиск успешно нашёл релевантные данные в базе совещаний:\n\n"
                    )
                    for idx, c in enumerate(citations[:4], 1):
                        fallback_text += f"{idx}. **{c['meeting_title']}** — *{c['source_label']}*:\n> {c['snippet']}\n\n"
                    fallback_text += (
                        "💡 *Для генерации полных ответов скачайте модель командой:*\n"
                        f"`ollama pull {settings.LLM_MODEL or 'qwen2.5:latest'}`\n"
                        "или укажите API-ключ в разделе **Настройки**."
                    )
                else:
                    fallback_text = (
                        "⚠️ **Не удалось связаться с языковой моделью.**\n\n"
                        "Убедитесь, что Ollama запущена (`ollama serve`) и модель загружена:\n"
                        f"`ollama pull {settings.LLM_MODEL or 'qwen2.5:latest'}`\n"
                        "Либо настройте параметры модели в **Настройках**."
                    )

                asst_msg_id = str(uuid.uuid4())
                with get_db() as conn:
                    conn.execute(
                        "INSERT INTO chat_messages (id, session_id, role, content, citations) VALUES (?, ?, ?, ?, ?)",
                        (user_msg_id, session_id, "user", content, None),
                    )
                    conn.execute(
                        "INSERT INTO chat_messages (id, session_id, role, content, citations) VALUES (?, ?, ?, ?, ?)",
                        (asst_msg_id, session_id, "assistant", fallback_text, json.dumps(citations, ensure_ascii=False))
                    )
                    conn.execute("UPDATE chat_sessions SET updated_at = CURRENT_TIMESTAMP WHERE id = ?", (session_id,))
                    conn.commit()

                yield f"event: delta\ndata: {json.dumps({'delta': fallback_text}, ensure_ascii=False)}\n\n"
                yield f"event: done\ndata: {json.dumps({'message_id': asst_msg_id, 'full_text': fallback_text}, ensure_ascii=False)}\n\n"
            else:
                yield f"event: error\ndata: {json.dumps({'error': str(e)})}\n\n"

    return StreamingResponse(event_generator(), media_type="text/event-stream")

@app.post("/api/v1/search/reindex")
async def reindex_search():
    with get_db() as conn:
        stats = await indexer.reindex_all_meetings(conn)
        return {"status": "ok", **stats}

@app.get("/api/v1/search/stats")
async def get_search_stats():
    with get_db() as conn:
        return vector_store.get_stats(conn)

@app.post("/api/v1/edit")
async def edit_text(req: EditRequest):
    sync_generation_settings()
    system_prompt = (
        "You are an editorial assistant for official and business meeting protocols. "
        "Your job is to rewrite or correct the highlighted text according to the user's instructions. "
        "Preserve official business style and language (Kazakh or Russian). Return ONLY the edited replacement text, "
        "with no explanation or greeting."
    )

    user_prompt = f"Original text:\n\"\"\"\n{req.selected_text}\n\"\"\"\n\nInstruction:\n{req.instruction}\n\nEdited text:"

    try:
        edited = await llm.complete_messages([
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ], temperature=0.2)
        return {"edited_text": edited.strip()}
    except llm.LLMError as error:
        raise HTTPException(502, str(error)) from None

# ---------------------------------------------------------------------------
# DOCX Export & File Actions
# ---------------------------------------------------------------------------

def _build_meeting_docx(
    meeting_id: str,
    lang: str = "ru",
    template_id: Optional[str] = None,
    mode: Optional[str] = None,
) -> tuple[bytes, str]:
    with get_db() as conn:
        row = conn.execute("SELECT * FROM meetings WHERE id = ?", (meeting_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Meeting not found")

    title = row["title"] or "Совещание"

    # If custom template requested or attached to meeting
    effective_tpl_id = template_id or (row["template_id"] if "template_id" in row.keys() else None)
    if mode != "standard" and effective_tpl_id:
        with get_db() as conn:
            tpl_row = conn.execute("SELECT * FROM protocol_templates WHERE id = ?", (effective_tpl_id,)).fetchone()
        if tpl_row and os.path.exists(tpl_row["docx_path"]):
            with open(tpl_row["docx_path"], "rb") as f:
                tpl_bytes = f.read()

            parsed_tpl = docx_template.parse_docx_template(tpl_bytes)
            slots_data = json.loads(tpl_row["slots"] or "[]")
            stored_slots_map = {s["key"]: s for s in slots_data}
            loop_iters = docx_template._loop_iterables(tpl_bytes)
            for s in parsed_tpl.slots:
                st = stored_slots_map.get(s.key)
                if st:
                    s.label = st.get("label", s.label)
                    s.value_type = st.get("value_type", s.value_type)
                    s.repeat = st.get("repeat", s.repeat)
                    s.omit_when_empty = st.get("omit_when_empty", s.omit_when_empty)
                elif s.key in loop_iters:
                    s.value_type = "list[object]"

            raw_proto = (row["protocol_kz"] if lang == "kz" else row["protocol_ru"]) or "{}"
            try:
                protocol_data = json.loads(raw_proto)
            except Exception:
                protocol_data = {}

            tpl_values = _adapt_meeting_to_template_values(parsed_tpl.slots, dict(row), protocol_data, lang=lang)
            try:
                rendered_bytes = docx_template.render_docx_template(tpl_bytes, tpl_values, slots=parsed_tpl.slots)
            except Exception as ren_e:
                logger.warning("Render template with adapted values failed (%s), retrying with mock fallback", ren_e)
                mock_vals = _mock_template_values(parsed_tpl.slots)
                tpl_values.update({k: v for k, v in mock_vals.items() if k not in tpl_values or not tpl_values[k]})
                rendered_bytes = docx_template.render_docx_template(tpl_bytes, tpl_values, slots=parsed_tpl.slots)

            return rendered_bytes, title

    # Standard docx generation fallback
    import docx
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = docx.Document()

    raw_protocol = (row["protocol_kz"] if lang == "kz" else row["protocol_ru"]) or "{}"
    try:
        protocol_data = json.loads(raw_protocol)
    except Exception:
        protocol_data = {}

    # Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(f"ХАТТАМА / ПРОТОКОЛ\n{title.upper()}")
    run.font.bold = True
    run.font.size = Pt(16)

    # Date
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    created_date = (row['created_at'] or "")[:10]
    date_p.add_run(f"Күні / Дата: {created_date}")

    # Agenda
    doc.add_heading("Повестка дня / Күн тәртібі", level=2)
    agenda_text = row["agenda"] or protocol_data.get("metadata", {}).get("agenda", "Вопросы рабочего совещания")
    doc.add_paragraph(agenda_text)

    # Participants
    doc.add_heading("Присутствовали / Қатысқандар", level=2)
    protocol_participants = protocol_data.get("participants")
    if protocol_participants and isinstance(protocol_participants, list):
        participants = protocol_participants
    else:
        try:
            participants = json.loads(row["participants"] or "[]")
        except Exception:
            participants = []

    if participants:
        for p in participants:
            if isinstance(p, dict):
                name = p.get("name", "Участник")
                pos = f" — {p.get('position')}" if p.get("position") else ""
                doc.add_paragraph(f"• {name}{pos}")
            else:
                doc.add_paragraph(f"• {p}")
    else:
        doc.add_paragraph("Согласно списку участников")

    # Decisions / Topics
    doc.add_heading("Решения и поручения / Шешімдер", level=2)
    raw_topics = protocol_data.get("agenda_items") or protocol_data.get("topics") or []
    if raw_topics:
        for idx, t in enumerate(raw_topics, 1):
            topic_title = t.get("topic") or t.get("topic_name") or f"Вопрос {idx}"
            doc.add_heading(f"{idx}. {topic_title}", level=3)

            speaker = t.get("speaker") or t.get("discussion")
            if speaker:
                doc.add_paragraph(f"Докладчик / Баяндамашы: {speaker}")

            decisions = t.get("decisions", [])
            for d in decisions:
                if isinstance(d, dict):
                    text = d.get("decision", "")
                    assignee = d.get("responsible", "")
                    deadline = d.get("deadline", "")
                    meta = []
                    if assignee:
                        meta.append(f"Отв: {assignee}")
                    if deadline:
                        meta.append(f"Срок: {deadline}")
                    meta_str = f" ({', '.join(meta)})" if meta else ""
                    doc.add_paragraph(f"— {text}{meta_str}")
                else:
                    doc.add_paragraph(f"— {d}")
    else:
        raw_summary = (row["summary_kz"] if lang == "kz" else row["summary_ru"]) or "{}"
        try:
            summary_data = json.loads(raw_summary)
        except Exception:
            summary_data = {}
        if summary_data.get("executive_summary"):
            doc.add_paragraph(summary_data["executive_summary"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), title


@app.get("/api/v1/meetings/{meeting_id}/export/docx")
async def export_docx(
    meeting_id: str,
    lang: str = Query("ru", pattern="^(ru|kz)$"),
    template_id: Optional[str] = Query(None),
    mode: Optional[str] = Query(None),
):
    docx_bytes, title = _build_meeting_docx(meeting_id, lang=lang, template_id=template_id, mode=mode)

    export_path = EXPORTS_DIR / f"{meeting_id}_{lang}.docx"
    export_path.write_bytes(docx_bytes)

    clean_title = re.sub(r'[/\\?%*:|"<>]+', '-', title).strip() or f"meeting_{meeting_id[:8]}"
    prefix = "Хаттама" if lang == "kz" else "Протокол"
    lang_tag = "KZ" if lang == "kz" else "RU"
    utf8_filename = f"{prefix} - {clean_title} ({lang_tag}).docx"
    ascii_fallback = f"Protocol_{meeting_id[:8]}_{lang}.docx"
    quoted_filename = urllib.parse.quote(utf8_filename)

    headers = {
        "Content-Disposition": f'attachment; filename="{ascii_fallback}"; filename*=UTF-8\'\'{quoted_filename}'
    }
    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers=headers,
    )


class SaveExportRequest(BaseModel):
    lang: str = "ru"
    template_id: Optional[str] = None
    open_folder: bool = True
    open_file: bool = False


@app.post("/api/v1/meetings/{meeting_id}/export/docx/save")
async def save_docx_to_downloads(
    meeting_id: str,
    req: SaveExportRequest = SaveExportRequest(),
):
    docx_bytes, title = _build_meeting_docx(meeting_id, lang=req.lang, template_id=req.template_id)

    # Save cached copy
    cached_path = EXPORTS_DIR / f"{meeting_id}_{req.lang}.docx"
    cached_path.write_bytes(docx_bytes)

    clean_title = re.sub(r'[/\\?%*:|"<>]+', '-', title).strip() or f"meeting_{meeting_id[:8]}"
    prefix = "Хаттама" if req.lang == "kz" else "Протокол"
    lang_tag = "KZ" if req.lang == "kz" else "RU"
    base_name = f"{prefix} - {clean_title} ({lang_tag})"

    downloads_dir = Path.home() / "Downloads"
    if not downloads_dir.exists():
        try:
            downloads_dir.mkdir(parents=True, exist_ok=True)
        except Exception:
            downloads_dir = EXPORTS_DIR

    final_filename = f"{base_name}.docx"
    target_path = downloads_dir / final_filename
    counter = 1
    while target_path.exists():
        final_filename = f"{base_name} ({counter}).docx"
        target_path = downloads_dir / final_filename
        counter += 1

    target_path.write_bytes(docx_bytes)

    if req.open_file:
        try:
            if sys.platform == "darwin":
                subprocess.Popen(["open", str(target_path)])
            elif sys.platform == "win32":
                os.startfile(str(target_path))
            else:
                subprocess.Popen(["xdg-open", str(target_path)])
        except Exception as e:
            logger.warning("Failed to open file: %s", e)
    elif req.open_folder:
        try:
            if sys.platform == "darwin":
                subprocess.Popen(["open", "-R", str(target_path)])
            elif sys.platform == "win32":
                subprocess.Popen(["explorer", f"/select,{str(target_path)}"])
            else:
                subprocess.Popen(["xdg-open", str(downloads_dir)])
        except Exception as e:
            logger.warning("Failed to reveal file in folder: %s", e)

    return {
        "success": True,
        "filename": final_filename,
        "path": str(target_path),
    }


@app.get("/api/v1/meetings/{meeting_id}/export/pdf")
async def export_pdf(
    meeting_id: str,
    lang: str = Query("ru", pattern="^(ru|kz)$"),
):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM meetings WHERE id = ?", (meeting_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Meeting not found")

    pdf_bytes, title = pdf_export.build_meeting_protocol_pdf(dict(row), lang=lang)

    export_path = EXPORTS_DIR / f"{meeting_id}_{lang}.pdf"
    export_path.write_bytes(pdf_bytes)

    clean_title = re.sub(r'[/\\?%*:|"<>]+', '-', title).strip() or f"meeting_{meeting_id[:8]}"
    prefix = "Хаттама" if lang == "kz" else "Протокол"
    lang_tag = "KZ" if lang == "kz" else "RU"
    utf8_filename = f"{prefix} - {clean_title} ({lang_tag}).pdf"
    ascii_fallback = f"Protocol_{meeting_id[:8]}_{lang}.pdf"
    quoted_filename = urllib.parse.quote(utf8_filename)

    headers = {
        "Content-Disposition": f'attachment; filename="{ascii_fallback}"; filename*=UTF-8\'\'{quoted_filename}'
    }
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers=headers,
    )


@app.post("/api/v1/meetings/{meeting_id}/export/pdf/save")
async def save_pdf_to_downloads(
    meeting_id: str,
    req: SaveExportRequest = SaveExportRequest(),
):
    with get_db() as conn:
        row = conn.execute("SELECT * FROM meetings WHERE id = ?", (meeting_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Meeting not found")

    pdf_bytes, title = pdf_export.build_meeting_protocol_pdf(dict(row), lang=req.lang)

    # Save cached copy
    cached_path = EXPORTS_DIR / f"{meeting_id}_{req.lang}.pdf"
    cached_path.write_bytes(pdf_bytes)

    clean_title = re.sub(r'[/\\?%*:|"<>]+', '-', title).strip() or f"meeting_{meeting_id[:8]}"
    prefix = "Хаттама" if req.lang == "kz" else "Протокол"
    lang_tag = "KZ" if req.lang == "kz" else "RU"
    base_name = f"{prefix} - {clean_title} ({lang_tag})"

    downloads_dir = Path.home() / "Downloads"
    if not downloads_dir.exists():
        try:
            downloads_dir.mkdir(parents=True, exist_ok=True)
        except Exception:
            downloads_dir = EXPORTS_DIR

    final_filename = f"{base_name}.pdf"
    target_path = downloads_dir / final_filename
    counter = 1
    while target_path.exists():
        final_filename = f"{base_name} ({counter}).pdf"
        target_path = downloads_dir / final_filename
        counter += 1

    target_path.write_bytes(pdf_bytes)

    if req.open_file:
        try:
            if sys.platform == "darwin":
                subprocess.Popen(["open", str(target_path)])
            elif sys.platform == "win32":
                os.startfile(str(target_path))
            else:
                subprocess.Popen(["xdg-open", str(target_path)])
        except Exception as e:
            logger.warning("Failed to open PDF file: %s", e)
    elif req.open_folder:
        try:
            if sys.platform == "darwin":
                subprocess.Popen(["open", "-R", str(target_path)])
            elif sys.platform == "win32":
                subprocess.Popen(["explorer", f"/select,{str(target_path)}"])
            else:
                subprocess.Popen(["xdg-open", str(downloads_dir)])
        except Exception as e:
            logger.warning("Failed to reveal PDF in folder: %s", e)

    return {
        "success": True,
        "filename": final_filename,
        "path": str(target_path),
    }


class SystemFileActionRequest(BaseModel):
    path: str


@app.post("/api/v1/system/open-file")
async def system_open_file(req: SystemFileActionRequest):
    file_path = Path(req.path).expanduser().resolve()
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    try:
        if sys.platform == "darwin":
            subprocess.Popen(["open", str(file_path)])
        elif sys.platform == "win32":
            os.startfile(str(file_path))
        else:
            subprocess.Popen(["xdg-open", str(file_path)])
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to open file: {e}")
    return {"success": True}


@app.post("/api/v1/system/reveal-file")
async def system_reveal_file(req: SystemFileActionRequest):
    file_path = Path(req.path).expanduser().resolve()
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    try:
        if sys.platform == "darwin":
            subprocess.Popen(["open", "-R", str(file_path)])
        elif sys.platform == "win32":
            subprocess.Popen(["explorer", f"/select,{str(file_path)}"])
        else:
            subprocess.Popen(["xdg-open", str(file_path.parent)])
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to reveal file: {e}")
    return {"success": True}


class SaveFileUrlRequest(BaseModel):
    url: str
    filename: str
    open_folder: bool = True


@app.post("/api/v1/system/save-url-to-downloads")
async def system_save_url_to_downloads(req: SaveFileUrlRequest):
    downloads_dir = Path.home() / "Downloads"
    downloads_dir.mkdir(parents=True, exist_ok=True)

    clean_filename = re.sub(r'[/\\?%*:|"<>]+', '-', req.filename).strip()
    target_path = downloads_dir / clean_filename
    base, ext = os.path.splitext(clean_filename)
    counter = 1
    while target_path.exists():
        target_path = downloads_dir / f"{base} ({counter}){ext}"
        counter += 1

    async with httpx.AsyncClient() as client:
        resp = await client.get(req.url, follow_redirects=True)
        if resp.status_code != 200:
            raise HTTPException(status_code=resp.status_code, detail="Failed to fetch file content")
        target_path.write_bytes(resp.content)

    if req.open_folder:
        try:
            if sys.platform == "darwin":
                subprocess.Popen(["open", "-R", str(target_path)])
            elif sys.platform == "win32":
                subprocess.Popen(["explorer", f"/select,{str(target_path)}"])
            else:
                subprocess.Popen(["xdg-open", str(downloads_dir)])
        except Exception as e:
            logger.warning("Failed to reveal file: %s", e)

    return {
        "success": True,
        "filename": target_path.name,
        "path": str(target_path),
    }

# ---------------------------------------------------------------------------
# Diagnostics & System Endpoints
# ---------------------------------------------------------------------------

@app.get("/api/v1/system/health")
async def backend_health():
    # Readiness is independent of Ollama, STT dependencies and downloaded models.
    return {"service": "steppe-meeting", "backend_dir": str(BASE_DIR), "pid": os.getpid(),
            "app_version": os.environ.get("STEPPE_APP_VERSION", "development"),
            "runtime": "bundled" if os.environ.get("STEPPE_BUNDLED_RUNTIME") == "1" else "development",
            "update_receipt": os.environ.get("STEPPE_UPDATE_RECEIPT", "")}


@app.get("/api/v1/system/status")
async def system_status():
    llm_url = get_setting("llm_base_url", "http://localhost:11434/v1").rstrip("/")
    api_key = get_setting("llm_api_key", "").strip()
    current_model = get_setting("llm_model", "qwen2.5:latest").strip()
    whisper_url = get_setting("whisper_base_url", "http://localhost:8000/v1").rstrip("/")

    headers = get_auth_headers(api_key)
    tags_url = get_ollama_tags_url(llm_url)
    models_url = get_models_url(llm_url)

    ollama_ok = False
    ollama_err = ""
    ollama_models = []
    is_cloud_provider = any(domain in llm_url for domain in ("api.openai.com", "openrouter.ai", "api.groq.com", "api.deepseek.com"))

    cli_status = None
    llm_mode = get_setting("llm_mode", "api")
    if llm_mode == "cli":
        provider = get_setting("llm_cli_provider", "opencode")
        options = get_cli_config().get(provider, {})
        cli_status = await cli_llm.check(provider, options.get("path", ""), cached=True)
        ollama_err = cli_status.get("error", "")
        ollama_ok = cli_status["available"]
        model_ready = cli_llm.model_ready(cli_status, options.get("model", ""))
        access = cli_status.get("model_access") or {}
        if not ollama_err and access.get("model") == options.get("model", "").strip() and not access.get("ok"):
            ollama_err = access.get("error", "")
        llm_url = cli_status.get("path", provider)
        current_model = options.get("model") or "CLI default"
    else:
        try:
            async with httpx.AsyncClient(timeout=3.5) as client:
                resp = await client.get(tags_url)
                if resp.status_code == 200:
                    ollama_ok = True
                    data = resp.json()
                    ollama_models = [m.get("name") for m in data.get("models", []) if m.get("name")]
                else:
                    resp2 = await client.get(models_url, headers=headers)
                    if resp2.status_code == 200:
                        ollama_ok = True
                        data = resp2.json()
                        ollama_models = [m.get("id") for m in data.get("data", []) if m.get("id")]
                    elif resp2.status_code in (401, 403):
                        ollama_err = f"Ошибка авторизации ({resp2.status_code}): проверьте Bearer-токен / API ключ."
                    else:
                        ollama_err = f"Эндпоинт вернул код {resp2.status_code}"
        except Exception as e:
            ollama_err = str(e)

        model_ready = False
        if ollama_ok:
            if is_cloud_provider:
                model_ready = bool(api_key)
                if not model_ready:
                    ollama_err = "Не указан API-ключ для внешнего провайдера"
            else:
                if ollama_models:
                    for m in ollama_models:
                        # Exact match, or base match without :latest / tag
                        if m == current_model or m.split(":")[0] == current_model.split(":")[0]:
                            model_ready = True
                            break
                if not model_ready:
                    if not ollama_models:
                        ollama_err = f"Ollama запущена, но нет скачанных моделей. Выполните: `ollama pull {current_model}`"
                    else:
                        ollama_err = f"Модель '{current_model}' не скачана в Ollama. Выполните: `ollama pull {current_model}` (доступные: {', '.join(ollama_models)})"

    whisper_mode = get_setting("whisper_mode", "local")
    whisper_url = get_setting("whisper_base_url", "http://localhost:8000/v1").rstrip("/")
    whisper_key = get_setting("whisper_api_key", "").strip()
    whisper_local_model = get_setting("whisper_local_model", "small")
    whisper_ext_model = get_setting("whisper_model", "whisper-1")

    whisper_ok = False
    whisper_err = ""

    if whisper_mode == "local":
        local_status = whisper_service.get_whisper_full_status(
            whisper_mode, whisper_url, whisper_ext_model, whisper_key, whisper_local_model,
        )
        whisper_ok = local_status["ready"]
        whisper_err = local_status["local"]["error"] or ""
    else:
        auth_hdrs = get_auth_headers(whisper_key)
        try:
            async with httpx.AsyncClient(timeout=3.0) as client:
                resp = await client.get(f"{whisper_url}/models", headers=auth_hdrs)
                if resp.status_code == 200:
                    whisper_ok = True
                elif resp.status_code in (401, 403):
                    whisper_err = f"Ошибка авторизации ({resp.status_code}): проверьте Bearer-токен Whisper"
                else:
                    whisper_err = f"Эндпоинт вернул код {resp.status_code}"
        except Exception as e:
            whisper_err = str(e)

    whisper_status_obj = {
        "connected": whisper_ok,
        "mode": whisper_mode,
        "url": whisper_url,
        "model": whisper_local_model if whisper_mode == "local" else whisper_ext_model,
        "error": whisper_err,
        "local": whisper_service.get_whisper_full_status(
            active_mode=whisper_mode,
            external_url=whisper_url,
            external_model=whisper_ext_model,
            external_key=whisper_key,
            local_model_name=whisper_local_model,
        )["local"],
    }

    return {
        "llm_mode": llm_mode,
        **({"cli": cli_status} if cli_status else {}),
        "ollama": {
            "connected": ollama_ok and model_ready,
            "service_online": ollama_ok,
            "model_ready": model_ready,
            "url": llm_url,
            "model": current_model,
            "available_models": ollama_models,
            "error": ollama_err,
        },
        "whisper": whisper_status_obj,
    }

@app.get("/api/v1/system/whisper/status")
async def get_whisper_status():
    whisper_mode = get_setting("whisper_mode", "local")
    whisper_local_model = get_setting("whisper_local_model", "small")
    whisper_url = get_setting("whisper_base_url", "http://localhost:8000/v1").rstrip("/")
    whisper_model = get_setting("whisper_model", "whisper-1")
    whisper_key = get_setting("whisper_api_key", "").strip()
    return whisper_service.get_whisper_full_status(
        active_mode=whisper_mode,
        external_url=whisper_url,
        external_model=whisper_model,
        external_key=whisper_key,
        local_model_name=whisper_local_model,
    )

@app.post("/api/v1/system/whisper/install-env")
async def install_whisper_env(req: Optional[WhisperModelActionRequest] = None):
    try:
        model = req.model if req else get_setting("whisper_local_model", "small")
        return whisper_service.start_install_env(model)
    except (ValueError, RuntimeError) as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/v1/system/whisper/download-model")
async def download_whisper_model(req: WhisperModelActionRequest):
    try:
        return whisper_service.start_download_model(req.model)
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.post("/api/v1/system/whisper/delete-model")
async def delete_whisper_model(req: WhisperModelActionRequest):
    try:
        return whisper_service.delete_downloaded_model(req.model)
    except (ValueError, RuntimeError) as e:
        raise HTTPException(status_code=400, detail=str(e))

@app.get("/api/v1/system/models")
async def list_models():
    llm_url = get_setting("llm_base_url", "http://localhost:11434/v1").rstrip("/")
    api_key = get_setting("llm_api_key", "").strip()
    return await discover_api_models(llm_url, api_key)


class ModelDiscoveryRequest(BaseModel):
    base_url: str = Field(max_length=4096)
    api_key: str = Field(default="", max_length=16384)


@app.post("/api/v1/system/models")
async def preview_models(req: ModelDiscoveryRequest):
    """Discover from the current form without persisting incomplete settings."""
    url = req.base_url.strip().rstrip("/")
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise HTTPException(status_code=400, detail="Укажите полный HTTP или HTTPS URL провайдера")
    return await discover_api_models(url, req.api_key.strip())


async def discover_api_models(llm_url: str, api_key: str):
    headers = get_auth_headers(api_key)

    tags_url = get_ollama_tags_url(llm_url)
    models_url = get_models_url(llm_url)

    # 1. Try Ollama tags endpoint
    try:
        async with httpx.AsyncClient(timeout=3.0) as client:
            resp = await client.get(tags_url)
            if resp.status_code == 200:
                data = resp.json()
                models = [m.get("name") for m in data.get("models", []) if m.get("name")]
                return sorted(models)
    except Exception:
        pass

    # 2. Try standard OpenAI /models endpoint
    try:
        async with httpx.AsyncClient(timeout=4.0) as client:
            resp2 = await client.get(models_url, headers=headers)
            if resp2.status_code == 200:
                data = resp2.json()
                models = [m.get("id") for m in data.get("data", []) if m.get("id")]
                if models:
                    return sorted(models)
    except Exception as e:
        logger.warning("Failed to list models from %s: %s", models_url, e)

    return []

class PullModelRequest(BaseModel):
    model_name: str

@app.post("/api/v1/system/pull-model")
async def pull_model(req: PullModelRequest, background_tasks: BackgroundTasks):
    llm_url = get_setting("llm_base_url", "http://localhost:11434/v1").rstrip("/")
    base_url = llm_url[:-3] if llm_url.endswith("/v1") else llm_url
    model_to_pull = req.model_name.strip()
    if not model_to_pull:
        raise HTTPException(status_code=400, detail="model_name is required")

    async def do_pull():
        try:
            logger.info("Starting background pull for model %s...", model_to_pull)
            async with httpx.AsyncClient(timeout=1800.0) as client:
                res = await client.post(f"{base_url}/api/pull", json={"name": model_to_pull, "stream": False})
                logger.info("Pull completed for %s: %s", model_to_pull, res.status_code)
        except Exception as e:
            logger.error("Failed to pull model %s: %s", model_to_pull, e)

    background_tasks.add_task(do_pull)
    return {"status": "pulling", "model": model_to_pull}

@app.get("/api/v1/system/config")
async def get_config():
    return {
        "update_channel": get_update_channel(),
        "llm_mode": get_setting("llm_mode", "api"),
        "llm_cli_provider": get_setting("llm_cli_provider", "opencode"),
        "llm_cli": get_cli_config(),
        "llm_base_url": get_setting("llm_base_url", "http://localhost:11434/v1"),
        "llm_model": get_setting("llm_model", "qwen2.5:latest"),
        "llm_api_key": get_setting("llm_api_key", ""),
        "whisper_mode": get_setting("whisper_mode", "local"),
        "whisper_local_model": get_setting("whisper_local_model", "small"),
        "whisper_device": get_setting("whisper_device", "auto"),
        "whisper_base_url": get_setting("whisper_base_url", "http://localhost:8000/v1"),
        "whisper_model": get_setting("whisper_model", "whisper-1"),
        "whisper_api_key": get_setting("whisper_api_key", ""),
    }

def get_update_channel() -> Literal["stable", "beta"]:
    channel = get_setting("update_channel", "stable")
    return channel if channel in ("stable", "beta") else "stable"


@app.get("/api/v1/system/updates")
async def get_updates(
    current_version: str = Query(..., max_length=200),
    channel: Optional[Literal["stable", "beta"]] = None,
    platform: str = Query("", max_length=32),
    arch: str = Query("", max_length=32),
    force: bool = False,
):
    result = await update_checker.check(current_version, channel or get_update_channel(), platform, arch, force)
    if platform in ("darwin", "macos") and arch in ("aarch64", "arm64"):
        return await updater.select(result)
    return result


@app.post("/api/v1/system/updates/channel")
async def save_update_channel(req: UpdateChannelRequest):
    # Independent from LLM/STT validation and saves: change exactly one key.
    set_setting("update_channel", req.channel)
    return {"channel": req.channel}


@app.post("/api/v1/system/config")
async def save_config(req: SystemConfigRequest):
    # Validate the entire proposed state before the transaction; inactive legacy
    # fields are retained and are not blockers when switching generation mode.
    current = await get_config()
    updates = req.model_dump(exclude_none=True)
    for key, value in list(updates.items()):
        if isinstance(value, str):
            updates[key] = value.strip()
    if "llm_cli" in updates:
        merged = current["llm_cli"].copy()
        for provider, options in updates["llm_cli"].items():
            merged[provider] = {**merged.get(provider, {}), **{k: v.strip() for k, v in options.items() if v is not None}}
        updates["llm_cli"] = merged
    proposed = {**current, **updates}
    try:
        if proposed["llm_mode"] not in ("api", "cli"):
            raise ValueError("Неизвестный способ генерации")
        if proposed["llm_cli_provider"] not in cli_llm.PROVIDERS:
            raise ValueError("Неизвестный CLI провайдер")
        if any(provider not in cli_llm.PROVIDERS for provider in proposed["llm_cli"]):
            raise ValueError("Неизвестный CLI провайдер")
        if proposed["llm_mode"] == "cli":
            provider = proposed["llm_cli_provider"]
            options = proposed["llm_cli"].get(provider, {})
            cli_llm.validate(provider, options.get("path", ""), options.get("model", ""), require_model=True)
        else:
            url = urllib.parse.urlparse(proposed["llm_base_url"])
            if url.scheme not in ("http", "https") or not url.netloc:
                raise ValueError("Укажите полный HTTP URL провайдера")
            if not proposed["llm_model"].strip():
                raise ValueError("Укажите модель для генерации")
        if proposed["whisper_mode"] == "local":
            whisper_service.get_model_info(proposed["whisper_local_model"])
    except ValueError as error:
        raise HTTPException(400, str(error)) from None
    with get_db() as conn:
        for key, value in updates.items():
            conn.execute("INSERT OR REPLACE INTO settings (key, value) VALUES (?, ?)",
                         (key, json.dumps(value, ensure_ascii=False) if key == "llm_cli" else value))
        conn.commit()
    sync_generation_settings()
    return {"success": True}


@app.post("/api/v1/system/cli/check")
async def check_cli(req: CliRequest):
    try:
        return await cli_llm.check(req.provider, (req.path or "").strip())
    except ValueError as error:
        raise HTTPException(400, str(error)) from None


@app.post("/api/v1/system/cli/models")
async def cli_models(req: CliRequest):
    return await cli_llm.models(req.provider, (req.path or "").strip())


@app.post("/api/v1/system/cli/test")
async def test_cli(req: CliTestRequest):
    return await cli_llm.test(req.provider, (req.path or "").strip(), (req.model or "").strip())


if __name__ == "__main__":
    port = int(os.environ.get("DESKTOP_API_PORT", 8008))
    logger.info("Starting Steppe Meeting Desktop Server on port %d...", port)
    uvicorn.run(app, host="127.0.0.1", port=port)
