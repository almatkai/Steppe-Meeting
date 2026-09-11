"""Parse and render user-supplied DOCX protocol templates.

DOCX is the source of layout. The application only extracts a data contract
from it and later fills Jinja variables with LLM-produced values; it does not
ask the model to recreate fonts, tables, headers, or spacing.

Explicit variables such as ``{{ employee_name }}`` are render-ready. Visual
blanks such as ``___`` and semantic instructions such as ``[employee name]``
are detected and normalized into variables in a separate working copy.
"""

from __future__ import annotations

import re
import zipfile
from copy import deepcopy
from dataclasses import asdict
from dataclasses import dataclass
from dataclasses import field
from io import BytesIO
from typing import Any

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import _Row
from docx.text.paragraph import Paragraph
from docxtpl import DocxTemplate
from jinja2 import StrictUndefined
from jinja2.sandbox import SandboxedEnvironment

# The original DOCX is never recompressed. Files above the preview threshold
# receive a separate optimized preview later; this parser still accepts them.
MAX_DOCX_BYTES = 50 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 250 * 1024 * 1024

PLACEHOLDER_RE = re.compile(r"\{\{\s*([a-zA-Z_][a-zA-Z0-9_]*)[^}]*\}\}")
SEMANTIC_MARKER_RE = re.compile(
    r"\[\s*((?=[^\[\]\r\n]*[^\W\d_])[^\[\]\r\n]{1,200}?)\s*\]"
)
VISUAL_MARKER_RE = re.compile(
    r"(?:_{3,}|(?<=20)_{2}|—{3,}|–{3,}|\[\s*_{2,}\s*\]|«\s*_{2,}\s*»|"
    r"\[\s*(?=[^\[\]\r\n]*[^\W\d_])[^\[\]\r\n]{1,200}?\s*\])"
)
VISUAL_FILL_RE = re.compile(r"(?:_{2,}|—{3,}|–{3,})")
PARENTHETICAL_CANDIDATE_RE = re.compile(
    r"\((?=[^()\r\n]{0,159}[^\W\d_])[^()\r\n]{1,160}\)"
)
_KAZAKH_MONTH = (
    r'қаңтар|ақпан|наурыз|сәуір|мамыр|маусым|шілде|тамыз|'
    r'қыркүйек|қазан|қараша|желтоқсан'
)
KAZAKH_DATE_PHRASE_RE = re.compile(
    rf'(?:20\s*_{{2,}}|_{{4,}})\s*жылғы\s*'
    rf'(?:'
    rf'«\s*_{{2,}}\s*»\s*_{{2,}}'
    rf'|'
    rf'_{{2,}}\s*(?:{_KAZAKH_MONTH})[а-яёәғқңөұүһі]*'
    rf'(?:\s+бен\s+_{{2,}}\s*(?:{_KAZAKH_MONTH})[а-яёәғқңөұүһі]*)?'
    rf')',
    re.IGNORECASE,
)


@dataclass
class Slot:
    key: str
    label: str
    value_type: str
    required: bool = False
    source: str = "placeholder"
    locations: list[str] = field(default_factory=list)
    raw_markers: list[str] = field(default_factory=list)
    repeat: dict[str, Any] | None = None
    omit_when_empty: bool = False
    inline: bool = False
    marker_contexts: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class ParsedTemplate:
    slots: list[Slot] = field(default_factory=list)
    paragraphs: list[dict[str, Any]] = field(default_factory=list)
    tables: list[dict[str, Any]] = field(default_factory=list)
    headers: list[str] = field(default_factory=list)
    footers: list[str] = field(default_factory=list)
    style_config: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    @property
    def render_ready(self) -> bool:
        return bool(self.slots) and all(slot.source == "placeholder" for slot in self.slots)


def _validate_docx(docx_bytes: bytes) -> None:
    if not docx_bytes:
        raise ValueError("DOCX is empty")
    if len(docx_bytes) > MAX_DOCX_BYTES:
        raise ValueError(f"DOCX exceeds the {MAX_DOCX_BYTES // 1024 // 1024} MB limit")

    try:
        with zipfile.ZipFile(BytesIO(docx_bytes)) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                raise ValueError("File is a ZIP archive, but not a DOCX document")
            uncompressed_size = sum(member.file_size for member in archive.infolist())
            if uncompressed_size > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("Uncompressed DOCX content is too large")
    except zipfile.BadZipFile as error:
        raise ValueError("File is not a valid DOCX document") from error


def validate_docx_file(docx_bytes: bytes) -> None:
    """Validate a DOCX container without requiring it to contain template fields."""
    _validate_docx(docx_bytes)


def _infer_type(key: str) -> str:
    folded = key.lower()
    words = set(re.findall(r'[a-zа-яёәғқңөұүһі]+', folded))
    if words.intersection({"date", "дата", "күні", "deadline"}):
        return "date"
    if any(
        token in folded
        for token in (
            "decisions",
            "agreements",
            "goals",
            "tasks",
            "participants",
            "candidates",
            "кандидат",
            "үміткер",
            "решения",
            "шешімдер",
            "сильные стороны",
            "зоны риска",
            "следующие шаги",
            "strengths",
            "risks",
            "next steps",
        )
    ):
        return "list[string]"
    if any(
        token in folded
        for token in ("summary", "description", "feedback", "notes", "заметки", "комментарий")
    ):
        return "text"
    return "string"


_CYRILLIC_TO_LATIN = str.maketrans({
    'а': 'a', 'б': 'b', 'в': 'v', 'г': 'g', 'д': 'd', 'е': 'e', 'ё': 'e',
    'ж': 'zh', 'з': 'z', 'и': 'i', 'й': 'i', 'к': 'k', 'л': 'l', 'м': 'm',
    'н': 'n', 'о': 'o', 'п': 'p', 'р': 'r', 'с': 's', 'т': 't', 'у': 'u',
    'ф': 'f', 'х': 'h', 'ц': 'c', 'ч': 'ch', 'ш': 'sh', 'щ': 'sh', 'ъ': '',
    'ы': 'y', 'ь': '', 'э': 'e', 'ю': 'yu', 'я': 'ya', 'ә': 'a', 'ғ': 'g',
    'қ': 'k', 'ң': 'n', 'ө': 'o', 'ұ': 'u', 'ү': 'u', 'һ': 'h', 'і': 'i',
})


def _slot_key(label: str, slots: dict[str, Slot], fallback: str) -> str:
    latin = label.lower().translate(_CYRILLIC_TO_LATIN)
    base = re.sub(r'[^a-z0-9]+', '_', latin).strip('_')[:60] or fallback
    if base[0].isdigit():
        base = f'field_{base}'[:60]
    key = base
    suffix = 2
    while key in slots:
        key = f'{base}_{suffix}'
        suffix += 1
    return key


def _label_from_context(text: str, marker_start: int, fallback: str) -> str:
    prefix = text[:marker_start].strip().rstrip(':;')
    if prefix:
        label = re.split(r"[\n\t:;]", prefix)[-1].strip(" №.-–—")
        if label:
            return label[-80:]
    return fallback


def _add_slot(
    slots: dict[str, Slot],
    *,
    key: str,
    label: str,
    source: str,
    location: str,
    raw: str,
    value_type: str | None = None,
    repeat: dict[str, Any] | None = None,
) -> None:
    existing = slots.get(key)
    if existing is None:
        existing = Slot(
            key=key,
            label=label,
            value_type=value_type or _infer_type(key),
            source=source,
            repeat=repeat,
        )
        slots[key] = existing
    if location not in existing.locations:
        existing.locations.append(location)
    if raw not in existing.raw_markers:
        existing.raw_markers.append(raw)
    if repeat is not None:
        existing.repeat = repeat


def _visual_matches(text: str):
    placeholder_spans = [match.span() for match in PLACEHOLDER_RE.finditer(text)]
    for match in VISUAL_MARKER_RE.finditer(text):
        if not any(match.start() < end and match.end() > start for start, end in placeholder_spans):
            yield match


def _semantic_marker_label(marker: str) -> str | None:
    match = SEMANTIC_MARKER_RE.fullmatch(marker)
    return match.group(1).strip() if match is not None else None


def _parenthetical_matches(text: str):
    protected_spans = [
        match.span()
        for pattern in (PLACEHOLDER_RE, SEMANTIC_MARKER_RE)
        for match in pattern.finditer(text)
    ]
    for match in PARENTHETICAL_CANDIDATE_RE.finditer(text):
        if any(match.start() < end and match.end() > start
               for start, end in protected_spans):
            continue
        if any(character in match.group(0) for character in '[]{}'):
            continue
        yield match


def parenthetical_candidates(parsed: ParsedTemplate) -> list[dict[str, Any]]:
    """Return round-bracket phrases whose field semantics require classification.

    Round brackets are normal prose too often to be promoted deterministically.
    Candidates therefore remain static until a separate semantic classifier
    explicitly selects them.
    """
    candidates = []
    paragraphs = parsed.paragraphs
    for index, paragraph in enumerate(paragraphs):
        text = paragraph['text']
        location = paragraph['location']
        neighbors = [
            {
                'location': item['location'],
                'text': item['text'][:240],
            }
            for item in paragraphs[max(0, index - 2):index + 3]
        ]
        for match in _parenthetical_matches(text):
            candidates.append({
                'id': f'{location}:{match.start()}:{match.end()}',
                'location': location,
                'start': match.start(),
                'end': match.end(),
                'text': match.group(0),
                'marked_text': (
                    text[:match.start()] + '⟦КАНДИДАТ⟧' + text[match.end():]
                ),
                'context': neighbors,
            })

    for table in parsed.tables:
        table_index = int(table['location'][len('body.tables['):-1])
        rows = table['rows']
        for row_index, row in enumerate(rows):
            for cell_index, text in enumerate(row):
                location = (
                    f'body.tables[{table_index}].rows[{row_index}]'
                    f'.cells[{cell_index}]'
                )
                for match in _parenthetical_matches(text):
                    candidates.append({
                        'id': f'{location}:{match.start()}:{match.end()}',
                        'location': location,
                        'start': match.start(),
                        'end': match.end(),
                        'text': match.group(0),
                        'marked_text': (
                            text[:match.start()] + '⟦КАНДИДАТ⟧' + text[match.end():]
                        ),
                        'context': {
                            'rows': [
                                [cell[:240] for cell in cells]
                                for cells in rows[max(0, row_index - 2):row_index + 3]
                            ],
                        },
                    })
    return candidates


def promote_parenthetical_fields(docx_bytes: bytes, fields: list[dict]) -> bytes:
    """Replace classifier-approved parenthetical phrases with visual fields.

    The surrounding round brackets stay in the document, so the value renders
    inline as ``(value1, value2)`` instead of a vertical list. Example:
    ``(по списку)`` becomes ``([Присутствовавшие участники])`` and finally
    ``({{ prisutstvovavshie_uchastniki }})``.
    """
    if not fields:
        return docx_bytes
    parsed = parse_docx_template(docx_bytes)
    candidates = {candidate['id']: candidate
                  for candidate in parenthetical_candidates(parsed)}
    replacements_by_location: dict[str, list[tuple[int, int, str]]] = {}
    for selected_field in fields:
        candidate = candidates.get(selected_field.get('id'))
        label = selected_field.get('label')
        if candidate is None:
            raise ValueError('Unknown parenthetical field candidate')
        if (
            not isinstance(label, str)
            or not label.strip()
            or len(label.strip()) > 120
            or any(character in label for character in '[]{}')
        ):
            raise ValueError('Invalid parenthetical field label')
        replacements_by_location.setdefault(candidate['location'], []).append((
            candidate['start'], candidate['end'], f'([{label.strip()}])',
        ))

    document = Document(BytesIO(docx_bytes))
    for location, replacements in replacements_by_location.items():
        paragraph_match = re.fullmatch(r'body\.paragraphs\[(\d+)\]', location)
        if paragraph_match:
            paragraph = document.paragraphs[int(paragraph_match.group(1))]
        else:
            cell_parts = _location_parts(location)
            if cell_parts is None:
                raise ValueError('Unsupported parenthetical field location')
            table_index, row_index, cell_index = cell_parts
            cell = document.tables[table_index].rows[row_index].cells[cell_index]
            paragraph_replacements: dict[int, list[tuple[int, int, str]]] = {}
            offset = 0
            for paragraph_index, paragraph in enumerate(cell.paragraphs):
                paragraph_end = offset + len(paragraph.text or '')
                for start, end, replacement in replacements:
                    if offset <= start and end <= paragraph_end:
                        paragraph_replacements.setdefault(paragraph_index, []).append((
                            start - offset, end - offset, replacement,
                        ))
                offset = paragraph_end + 1  # python-docx joins cell paragraphs with \n
            if sum(map(len, paragraph_replacements.values())) != len(replacements):
                raise ValueError('Ambiguous parenthetical field paragraph')
            for paragraph_index, local_replacements in paragraph_replacements.items():
                paragraph = cell.paragraphs[paragraph_index]
                _replace_text_spans_in_paragraph(
                    paragraph, paragraph.text or '', local_replacements,
                )
            continue
        _replace_text_spans_in_paragraph(paragraph, paragraph.text or '', replacements)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def apply_parenthetical_field_metadata(
    parsed: ParsedTemplate, fields: list[dict],
) -> None:
    """Restore classifier-selected labels and value types after DOCX parsing.

    Promoted fields render inline inside their preserved brackets, so a list
    becomes ``(A, B)`` rather than one item per line.
    """
    for selected_field in fields:
        raw = f"[{selected_field['label']}]"
        matches = [
            slot for slot in parsed.slots
            if selected_field['location'] in slot.locations and raw in slot.raw_markers
        ]
        if len(matches) != 1:
            raise ValueError(
                f"Could not bind parenthetical field: {selected_field['id']}"
            )
        matches[0].label = selected_field['label']
        matches[0].value_type = selected_field['value_type']
        matches[0].inline = True


def _numbered_visual_row(text: str) -> tuple[int, int, str] | None:
    """Return ``(major, minor, marker)`` for a row such as ``2.1. ____``."""
    matches = list(_visual_matches(text))
    if len(matches) != 1:
        return None
    marker = matches[0]
    if text[marker.end():].strip():
        return None
    prefix = text[:marker.start()].strip()
    number = re.fullmatch(r'(\d+)\.(\d+)\.?', prefix)
    if number is None:
        return None
    return int(number.group(1)), int(number.group(2)), marker.group(0)


def _numbered_paragraph_groups(paragraphs: list[Any]) -> dict[int, dict[str, Any]]:
    """Find numbered fill-line prototypes that should grow and shrink as lists."""
    rows = []
    for index, paragraph in enumerate(paragraphs):
        parsed = _numbered_visual_row(paragraph.text or '')
        if parsed is not None:
            major, minor, marker = parsed
            rows.append((index, major, minor, marker))

    contexts: dict[int, dict[str, Any]] = {}
    position = 0
    while position < len(rows):
        group = [rows[position]]
        position += 1
        while position < len(rows):
            previous = group[-1]
            candidate = rows[position]
            if (
                candidate[0] != previous[0] + 1
                or candidate[1] != previous[1]
                or candidate[2] != previous[2] + 1
            ):
                break
            group.append(candidate)
            position += 1

        if len(group) < 2:
            continue
        first_index, major, _, _ = group[0]
        key = f'repeated_list_{major}_{first_index}'
        section = _nearest_section_label(paragraphs, first_index)
        label = f'{section} — список' if section else f'Повторяемый список {major}'
        repeat = {'kind': 'numbered_paragraphs', 'major': major, 'original_count': len(group)}
        for index, _, _, marker in group:
            contexts[index] = {
                'key': key,
                'label': label,
                'marker': marker,
                'repeat': repeat,
            }
    return contexts


def _paragraph_numbering(paragraph: Paragraph):
    """Read direct or style-inherited Word numbering (absent from .text)."""
    properties = [paragraph._p.pPr]
    style = paragraph.style
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        properties.append(style.element.pPr)
        style = style.base_style
    num_id = None
    level = None
    for properties_element in properties:
        if properties_element is None or properties_element.numPr is None:
            continue
        numbering = properties_element.numPr
        if num_id is None and numbering.numId is not None:
            num_id = numbering.numId.val
        if level is None and numbering.ilvl is not None:
            level = numbering.ilvl.val
    if not num_id:
        return None
    level = level or 0
    root = paragraph.part.numbering_part.element
    definitions = root.xpath(f'./w:num[@w:numId="{num_id}"]/w:abstractNumId')
    if not definitions:
        return None
    abstract_id = definitions[0].get(qn('w:val'))
    levels = root.xpath(f'./w:num[@w:numId="{num_id}"]/w:lvlOverride[@w:ilvl="{level}"]/w:lvl')
    if not levels:
        levels = root.xpath(
            f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]/w:lvl[@w:ilvl="{level}"]'
        )
    if not levels:
        return None
    formats = levels[0].findall(qn('w:numFmt'))
    if not formats or formats[0].get(qn('w:val')) != 'decimal':
        return None
    return level, levels[0]


def _outline_row(paragraph: Paragraph):
    text = (paragraph.text or '').strip()
    literal = re.match(r'^(\d+(?:\.\d+)*)(?:\.|\))?\s*', text)
    if literal:
        level = literal.group(1).count('.')
        marker = text[literal.end():].strip()
    else:
        numbering = _paragraph_numbering(paragraph)
        level = numbering[0] if numbering is not None else None
        marker = text
    if level is None:
        return None
    if not (VISUAL_MARKER_RE.fullmatch(marker) or PLACEHOLDER_RE.fullmatch(marker)):
        return None
    return level, marker


def _outline_paragraph_groups(paragraphs: list[Any]) -> dict[int, dict[str, Any]]:
    """Recognize fillable multilevel examples, including legacy scalar fields."""
    groups = {}
    index = 0
    while index < len(paragraphs):
        start = index
        rows = []
        while index < len(paragraphs):
            row = _outline_row(paragraphs[index])
            if row is None:
                break
            if rows and paragraphs[index - 1]._p.getnext() is not paragraphs[index]._p:
                break  # A table or other body element separates these lists.
            rows.append((index, *row))
            index += 1
        if not rows:
            index += 1
            continue
        levels = [row[1] for row in rows]
        if levels[0] != 0 or len(rows) < 2:
            continue
        if any(current > previous + 1 for previous, current in zip(levels, levels[1:])):
            continue
        repeat = {
            'kind': 'numbered_outline',
            'original_count': len(rows),
            'levels': max(levels) + 1,
            'columns': [
                {'key': 'level', 'label': 'Уровень: 0 — раздел, 1 — подпункт'},
                {'key': 'text', 'label': 'Содержание пункта'},
            ],
            'prototype_levels': levels,
        }
        for row_index, _, marker in rows:
            groups[row_index] = {
                'key': f'repeated_outline_{start}',
                'label': 'Пункты документа',
                'marker': marker,
                'repeat': repeat,
            }
    return groups


def _scan_text(
    text: str,
    location: str,
    slots: dict[str, Slot],
    visual_counter: list[int],
    visual_line_context: tuple[str, int] | None = None,
) -> None:
    for match in PLACEHOLDER_RE.finditer(text):
        key = match.group(1)
        _add_slot(
            slots,
            key=key,
            label=key.replace("_", " ").capitalize(),
            source="placeholder",
            location=location,
            raw=match.group(0),
        )

    visual_matches = list(_visual_matches(text))
    placeholder_spans = [match.span() for match in PLACEHOLDER_RE.finditer(text)]
    kazakh_dates = [
        match for match in KAZAKH_DATE_PHRASE_RE.finditer(text)
        if not any(
            match.start() < end and match.end() > start
            for start, end in placeholder_spans
        )
    ]
    consumed_visuals = set()
    for kazakh_date in kazakh_dates:
        city_markers = [
            match for match in visual_matches
            if match.end() <= kazakh_date.start()
            and 'қаласы' in text[match.end():kazakh_date.start()].lower()
        ]
        if '«' in kazakh_date.group(0) and len(city_markers) == 1:
            city_marker = city_markers[0]
            city_key = _slot_key('Қала', slots, f'field_{visual_counter[0] + 1}')
            visual_counter[0] += 1
            _add_slot(
                slots,
                key=city_key,
                label='Қала',
                source='visual',
                location=location,
                raw=city_marker.group(0),
                value_type='string',
            )
            consumed_visuals.add(city_marker.span())
        date_label = (
            'Күндер аралығы'
            if re.search(r'\sбен\s', kazakh_date.group(0), re.IGNORECASE)
            else 'Отырыс күні'
        )
        date_key = _slot_key(
            date_label, slots, f'field_{visual_counter[0] + 1}',
        )
        visual_counter[0] += 1
        _add_slot(
            slots,
            key=date_key,
            label=date_label,
            source='visual',
            location=location,
            raw=kazakh_date.group(0),
            value_type='date',
        )
        for visual in visual_matches:
            if (
                visual.start() < kazakh_date.end()
                and visual.end() > kazakh_date.start()
            ):
                consumed_visuals.add(visual.span())

    visual_matches = [
        match for match in visual_matches
        if match.span() not in consumed_visuals
    ]
    date_line = (
        len(visual_matches) == 4
        and '«' in text
        and re.search(r'\b20\s*_{2,}', text) is not None
        and text.lstrip().lower().startswith('г.')
        and not kazakh_dates
    )
    date_labels = (
        ('Город', 'string'),
        ('Дата заседания — день', 'date'),
        ('Дата заседания — месяц', 'date'),
        ('Дата заседания — год', 'date'),
    ) if date_line else None

    first_marker_start = visual_matches[0].start() if visual_matches else 0
    shared_label = _label_from_context(text, first_marker_start, '')
    for marker_index, match in enumerate(visual_matches):
        visual_counter[0] += 1
        semantic_label = _semantic_marker_label(match.group(0))
        if semantic_label is not None:
            label = semantic_label
            value_type = _infer_type(label)
        elif date_labels is not None:
            label, value_type = date_labels[marker_index]
        elif re.search(r'\bот\s*$', text[:match.start()], re.IGNORECASE) and re.match(
            r'\s*(?:года|г\.)', text[match.end():], re.IGNORECASE,
        ):
            label, value_type = 'Дата документа', 'date'
        elif visual_line_context is not None and len(visual_matches) == 1:
            context_label, line_number = visual_line_context
            label = f'{context_label} — {line_number}'
            value_type = _infer_type(context_label)
        else:
            fallback = f"Поле {visual_counter[0]}"
            if len(visual_matches) == 1:
                label = _label_from_context(text, match.start(), fallback)
            elif shared_label:
                label = f'{shared_label} — поле {marker_index + 1}'
            else:
                label = fallback
            value_type = _infer_type(label)
        key = _slot_key(label, slots, f'field_{visual_counter[0]}')
        _add_slot(
            slots,
            key=key,
            label=label,
            source="visual",
            location=location,
            raw=match.group(0),
            value_type=value_type,
        )
        slots[key].marker_contexts.append({
            'location': location, 'start': match.start(), 'end': match.end(),
            'marked_text': text[:match.start()] + '⟦ПОЛЕ⟧' + text[match.end():],
        })


def _mark_optional_overflow_lines(
    paragraphs: list[Any],
    slots: dict[str, Slot],
) -> None:
    """Mark a spare line between two labeled fields for removal when unused."""
    slots_by_location = {
        location: slot
        for slot in slots.values()
        for location in slot.locations
    }
    for index in range(1, len(paragraphs) - 1):
        if not _visual_only_line(paragraphs[index].text or ''):
            continue
        previous = paragraphs[index - 1].text or ''
        following = paragraphs[index + 1].text or ''
        previous_matches = list(_visual_matches(previous))
        following_matches = list(_visual_matches(following))
        if (
            len(previous_matches) != 1
            or len(following_matches) != 1
            or ':' not in previous[:previous_matches[0].start()]
            or ':' not in following[:following_matches[0].start()]
        ):
            continue
        slot = slots_by_location.get(f'body.paragraphs[{index}]')
        if slot is not None:
            slot.omit_when_empty = True


def _normal_style(document: Any) -> dict[str, Any]:
    try:
        style = document.styles["Normal"]
        return {
            "normal_font": style.font.name,
            "normal_size_pt": style.font.size.pt if style.font.size else None,
        }
    except (KeyError, AttributeError):
        return {}


def _has_bottom_border(paragraph: Any) -> bool:
    properties = paragraph._p.pPr
    borders = properties.find(qn('w:pBdr')) if properties is not None else None
    bottom = borders.find(qn('w:bottom')) if borders is not None else None
    if bottom is None:
        return False
    return (bottom.get(qn('w:val')) or 'single') not in {'nil', 'none'}


def _clean_label(text: str) -> str:
    return re.sub(r'^\s*\d+\.\s*', '', text).strip().rstrip(':')


def _nearest_section_label(paragraphs: list[Any], start: int) -> str:
    for index in range(start - 1, max(-1, start - 6), -1):
        text = (paragraphs[index].text or '').strip()
        if not text or (text.startswith('(') and text.endswith(')')):
            continue
        return _clean_label(text)
    return ''


def _visual_only_line(text: str) -> bool:
    matches = list(_visual_matches(text))
    if len(matches) != 1:
        return False
    match = matches[0]
    remainder = text[:match.start()] + text[match.end():]
    return not remainder.strip()


def _visual_line_contexts(paragraphs: list[Any]) -> dict[int, tuple[str, int]]:
    """Associate consecutive fill lines with their preceding field label."""
    contexts: dict[int, tuple[str, int]] = {}
    index = 0
    while index < len(paragraphs):
        if not _visual_only_line(paragraphs[index].text or ''):
            index += 1
            continue

        start = index
        while index < len(paragraphs) and _visual_only_line(paragraphs[index].text or ''):
            index += 1
        if index - start < 2:
            continue

        label = _nearest_section_label(paragraphs, start)
        if not label:
            continue
        for line_number, paragraph_index in enumerate(range(start, index), start=1):
            contexts[paragraph_index] = (label, line_number)
    return contexts


def _infer_bordered_paragraph_slots(document: Any, slots: dict[str, Slot]) -> None:
    paragraphs = document.paragraphs
    index = 0
    while index < len(paragraphs):
        paragraph = paragraphs[index]
        text = (paragraph.text or '').strip()
        if text:
            index += 1
            continue
        if not _has_bottom_border(paragraph):
            index += 1
            continue

        start = index
        while index < len(paragraphs):
            candidate = paragraphs[index]
            if (candidate.text or '').strip() or not _has_bottom_border(candidate):
                break
            index += 1

        # A normalized multi-line field places its variable on the first
        # bordered paragraph and leaves the remaining ruled lines blank.
        if (
            start > 0
            and PLACEHOLDER_RE.search(paragraphs[start - 1].text or '')
            and _has_bottom_border(paragraphs[start - 1])
        ):
            continue

        # A ruled spacer immediately above an explicitly labelled agenda topic
        # is decoration, not a second field to duplicate that topic into.
        if index < len(paragraphs):
            following = paragraphs[index].text or ''
            heading = _nearest_section_label(paragraphs, start).casefold()
            if (
                any(word in heading for word in ('повест', 'agenda', 'күн тәртібі'))
                and (SEMANTIC_MARKER_RE.fullmatch(following.strip())
                     or PLACEHOLDER_RE.fullmatch(following.strip()))
                and _has_bottom_border(paragraphs[index])
            ):
                continue

        label = _nearest_section_label(paragraphs, start) or f'Текстовый блок {start + 1}'
        key = _slot_key(label, slots, f'section_{start + 1}')
        _add_slot(
            slots,
            key=key,
            label=label,
            source='structural',
            location=f'body.paragraphs[{start}]',
            raw=f'bordered-paragraphs:{start}-{index - 1}',
            value_type='text',
        )


def _nearest_left_label(cells: list[Any], cell_index: int) -> str:
    for index in range(cell_index - 1, -1, -1):
        text = (cells[index].text or '').strip()
        if text:
            return _clean_label(text)
    return ''


def _preceding_table_label(table: Any) -> str:
    element = table._tbl.getprevious()
    while element is not None:
        if element.tag == qn('w:p'):
            text = ''.join(node.text or '' for node in element.iter(qn('w:t'))).strip()
            if text:
                return _clean_label(text)
        element = element.getprevious()
    return ''


def _numbered_table_groups(document: Any) -> dict[tuple[int, int], dict[str, Any]]:
    """Find numbered prototype rows whose records may outgrow the source table."""
    contexts: dict[tuple[int, int], dict[str, Any]] = {}
    for table_index, table in enumerate(document.tables):
        numbered_rows = []
        for row_index, row in enumerate(table.rows):
            texts = [(cell.text or '').strip() for cell in row.cells]
            if not texts or re.fullmatch(r'\d+\.', texts[0]) is None:
                continue
            if any(text and not _visual_only_line(text) for text in texts[1:]):
                continue
            numbered_rows.append(row_index)

        position = 0
        while position < len(numbered_rows):
            indices = [numbered_rows[position]]
            position += 1
            while (
                position < len(numbered_rows)
                and numbered_rows[position] == indices[-1] + 1
            ):
                indices.append(numbered_rows[position])
                position += 1
            if len(indices) < 2:
                continue

            start_row = indices[0]
            header_row = table.rows[start_row - 1] if start_row > 0 else None
            used_column_keys: dict[str, Slot] = {}
            columns = []
            for cell_index in range(1, len(table.columns)):
                header = (
                    _clean_label(header_row.cells[cell_index].text or '')
                    if header_row is not None
                    else ''
                )
                label = header or f'Колонка {cell_index + 1}'
                column_key = _slot_key(
                    label, used_column_keys, f'column_{cell_index}',
                )
                used_column_keys[column_key] = Slot(
                    key=column_key,
                    label=label,
                    value_type='string',
                )
                columns.append({
                    'cell': cell_index,
                    'key': column_key,
                    'label': label,
                })

            repeat = {
                'kind': 'table_rows',
                'table': table_index,
                'start_row': start_row,
                'columns': columns,
                'original_count': len(indices),
            }
            group = {
                'key': f'repeated_table_{table_index}_{start_row}',
                'label': _preceding_table_label(table) or 'Повторяемая таблица',
                'repeat': repeat,
            }
            for row_index in indices:
                contexts[(table_index, row_index)] = group
    return contexts


def _infer_table_slots(
    document: Any,
    slots: dict[str, Slot],
    repeated_rows: set[tuple[int, int]],
) -> None:
    for table_index, table in enumerate(document.tables):
        if all((table_index, row_index) in repeated_rows for row_index in range(len(table.rows))):
            continue
        table_label = _preceding_table_label(table)
        header_cells = table.rows[0].cells
        for row_index, row in enumerate(table.rows):
            if (table_index, row_index) in repeated_rows:
                continue
            texts = [(cell.text or '').strip() for cell in row.cells]
            if any(re.search(r'\{%\s*tr\b', text) for text in texts):
                continue
            if any(texts) and all(text.startswith('☐') for text in texts if text):
                if all(texts) and not any(PLACEHOLDER_RE.search(text) for text in texts):
                    label = 'Итоговое решение'
                    key = _slot_key(label, slots, f'table_{table_index}_choice')
                    _add_slot(
                        slots,
                        key=key,
                        label=label,
                        source='structural',
                        location=f'body.tables[{table_index}].rows[{row_index}].cells[0]',
                        raw='checkbox-choice:' + '|'.join(texts),
                        value_type='string',
                    )
                continue

            for cell_index, cell in enumerate(row.cells):
                if (cell.text or '').strip():
                    continue
                row_label = _nearest_left_label(row.cells, cell_index)
                column_label = ''
                if row_index > 0 and cell_index < len(header_cells):
                    column_label = _clean_label(header_cells[cell_index].text or '')
                if not row_label and not column_label:
                    continue

                if 'голосован' in table_label.lower() and column_label:
                    choice = column_label.strip('«»" ')
                    label = f'Голосов — {choice}'
                    key_label = label
                elif table_label and row_label and column_label:
                    label = f'{table_label} — участник {row_label} — {column_label}'
                    key_label = f'{row_label} — {column_label}'
                elif table_label and column_label:
                    label = f'{table_label} — {column_label}'
                    key_label = label
                else:
                    label = row_label
                    if column_label and column_label != row_label:
                        label = f'{row_label} — {column_label}'
                    key_label = label
                key = _slot_key(
                    key_label,
                    slots,
                    f'table_{table_index}_{row_index}_{cell_index}',
                )
                _add_slot(
                    slots,
                    key=key,
                    label=label,
                    source='structural',
                    location=(
                        f'body.tables[{table_index}].rows[{row_index}].cells[{cell_index}]'
                    ),
                    raw='empty-table-cell',
                    value_type=_infer_type(label),
                )


def parse_docx_template(
    docx_bytes: bytes,
    *,
    infer_structural: bool = True,
) -> ParsedTemplate:
    """Return a reviewable descriptor of a DOCX template.

    A template containing only explicit Jinja variables is ``render_ready``.
    Visual blanks and square-bracket instructions are emitted as ``visual``
    slots so they can be mapped to stable field names before publishing.
    """
    _validate_docx(docx_bytes)
    try:
        document = Document(BytesIO(docx_bytes))
    except (PackageNotFoundError, ValueError, KeyError) as error:
        raise ValueError("DOCX structure cannot be read") from error

    result = ParsedTemplate(style_config=_normal_style(document))
    slots: dict[str, Slot] = {}
    visual_counter = [0]
    visual_line_contexts = _visual_line_contexts(document.paragraphs)
    numbered_groups = _numbered_paragraph_groups(document.paragraphs)
    numbered_groups.update(_outline_paragraph_groups(document.paragraphs))
    numbered_table_groups = _numbered_table_groups(document)

    for paragraph_index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text or ""
        location = f"body.paragraphs[{paragraph_index}]"
        result.paragraphs.append(
            {
                "location": location,
                "text": text,
                "style": paragraph.style.name if paragraph.style else None,
                "alignment": (
                    str(paragraph.alignment) if paragraph.alignment is not None else None
                ),
            }
        )
        numbered_group = numbered_groups.get(paragraph_index)
        if numbered_group is not None:
            visual_counter[0] += 1
            _add_slot(
                slots,
                key=numbered_group['key'],
                label=numbered_group['label'],
                source='visual',
                location=location,
                raw=numbered_group['marker'],
                value_type=(
                    'list[object]' if numbered_group['repeat']['kind'] == 'numbered_outline'
                    else 'list[string]'
                ),
                repeat=numbered_group['repeat'],
            )
        else:
            _scan_text(
                text,
                location,
                slots,
                visual_counter,
                visual_line_contexts.get(paragraph_index),
            )

    _mark_optional_overflow_lines(document.paragraphs, slots)

    for table_index, table in enumerate(document.tables):
        rows: list[list[str]] = []
        for row_index, row in enumerate(table.rows):
            row_text: list[str] = []
            numbered_table_group = numbered_table_groups.get((table_index, row_index))
            for cell_index, cell in enumerate(row.cells):
                text = cell.text or ""
                row_text.append(text)
                if numbered_table_group is None:
                    _scan_text(
                        text,
                        f"body.tables[{table_index}].rows[{row_index}].cells[{cell_index}]",
                        slots,
                        visual_counter,
                    )
            if numbered_table_group is not None:
                _add_slot(
                    slots,
                    key=numbered_table_group['key'],
                    label=numbered_table_group['label'],
                    source='structural',
                    location=f'body.tables[{table_index}].rows[{row_index}]',
                    raw='numbered-table-rows',
                    value_type='list[object]',
                    repeat=numbered_table_group['repeat'],
                )
            rows.append(row_text)
        result.tables.append({"location": f"body.tables[{table_index}]", "rows": rows})

    for section_index, section in enumerate(document.sections):
        for kind, container, collected in (
            ("header", section.header, result.headers),
            ("footer", section.footer, result.footers),
        ):
            texts: list[str] = []
            for paragraph_index, paragraph in enumerate(container.paragraphs):
                text = paragraph.text or ""
                texts.append(text)
                _scan_text(
                    text,
                    f"sections[{section_index}].{kind}.paragraphs[{paragraph_index}]",
                    slots,
                    visual_counter,
                )
            collected.append("\n".join(texts))

    if infer_structural:
        _infer_bordered_paragraph_slots(document, slots)
        _infer_table_slots(document, slots, set(numbered_table_groups))

    # Jinja variables used only in control tags (for example `{% for item in
    # decisions %}`) are not caught by PLACEHOLDER_RE. docxtpl can discover
    # them from the actual Word XML, including headers and footers.
    jinja_environment = SandboxedEnvironment(undefined=StrictUndefined)
    try:
        undeclared = DocxTemplate(BytesIO(docx_bytes)).get_undeclared_template_variables(
            jinja_env=jinja_environment
        )
    except Exception as error:
        raise ValueError(f"DOCX contains invalid template syntax: {error}") from error
    # A direct `{{ item }}` inside a loop is a local loop variable, not input
    # data. Keep only names Jinja reports as undeclared inputs.
    for key, slot in list(slots.items()):
        if slot.source == "placeholder" and key not in undeclared:
            del slots[key]
    consumed_keys = {
        match.group(1)
        for group in numbered_groups.values()
        for match in PLACEHOLDER_RE.finditer(group['marker'])
    }
    for key in sorted(undeclared - consumed_keys):
        if key not in slots:
            _add_slot(
                slots,
                key=key,
                label=key.replace("_", " ").capitalize(),
                source="placeholder",
                location="jinja.control_tag",
                raw=key,
            )

    result.slots = list(slots.values())
    if not result.slots:
        result.warnings.append(
            "Поля не найдены. Добавьте переменные вида {{ topic }} в DOCX."
        )
    elif not result.render_ready:
        result.warnings.append(
            "Найдены визуальные или структурные пустые поля; рабочая копия должна быть нормализована."
        )
    return result


def _clear_run_properties(target: Any) -> None:
    target_properties = target._r.rPr
    if target_properties is not None:
        target._r.remove(target_properties)


def _copy_run_properties(source: Any, target: Any) -> None:
    _clear_run_properties(target)
    source_properties = source._r.rPr
    if source_properties is not None:
        target._r.insert(0, deepcopy(source_properties))


def _replace_text_spans_in_paragraph(
    paragraph: Any,
    text: str,
    replacements: list[tuple[int, int, str]],
) -> None:
    runs = list(paragraph.runs)
    if not runs:
        paragraph.add_run(text)
        runs = list(paragraph.runs)

    # Word may split one underline across several runs. Work backwards so
    # earlier offsets remain stable while preserving the first run's style.
    for start, end, replacement in reversed(replacements):
        offset = 0
        start_run = start_offset = end_run = end_offset = None
        for index, run in enumerate(runs):
            run_end = offset + len(run.text)
            if start_run is None and offset <= start < run_end:
                start_run, start_offset = index, start - offset
            if offset < end <= run_end:
                end_run, end_offset = index, end - offset
                break
            offset = run_end
        if start_run is None or end_run is None or start_offset is None or end_offset is None:
            continue
        original_span = text[start:end]
        begins_with_fill = VISUAL_FILL_RE.match(original_span) is not None
        if begins_with_fill and start_offset == 0:
            context_run = next(
                (
                    runs[index]
                    for index in range(start_run - 1, -1, -1)
                    if runs[index].text
                ),
                None,
            )
            if context_run is None:
                context_run = next(
                    (
                        runs[index]
                        for index in range(end_run + 1, len(runs))
                        if runs[index].text
                    ),
                    None,
                )
            if context_run is not None:
                _copy_run_properties(context_run, runs[start_run])
            else:
                _clear_run_properties(runs[start_run])
        if start_run == end_run:
            run = runs[start_run]
            run.text = run.text[:start_offset] + replacement + run.text[end_offset:]
            continue
        start_text = runs[start_run].text[:start_offset]
        end_text = runs[end_run].text[end_offset:]
        runs[start_run].text = start_text + replacement
        for index in range(start_run + 1, end_run):
            runs[index].text = ''
        runs[end_run].text = end_text


def _replace_visuals_in_paragraph(
    paragraph: Any,
    text: str,
    matches: list,
    slots_for_paragraph: list[Slot],
) -> None:
    replacements = []
    for (match, slot) in zip(matches, slots_for_paragraph):
        marker = match.group(0)
        if _semantic_marker_label(marker) is not None:
            replacements.append((match.start(), match.end(), f'{{{{ {slot.key} }}}}'))
            continue
        fill = VISUAL_FILL_RE.search(marker)
        if fill is None:
            continue
        expression = f'{{{{ {slot.key} or {fill.group(0)!r} }}}}'
        replacement = marker[:fill.start()] + expression + marker[fill.end():]
        if (
            fill.start() == 0
            and match.start() > 0
            and text[match.start() - 1].isalpha()
        ):
            replacement = ' ' + replacement
        if (
            fill.end() == len(marker)
            and match.end() < len(text)
            and text[match.end()].isalpha()
        ):
            replacement += ' '
        replacements.append((match.start(), match.end(), replacement))
    _replace_text_spans_in_paragraph(paragraph, text, replacements)


def _replace_visuals_in_cell(cell: Any, slots: list[Slot]) -> None:
    """Replace visual markers inside one table cell, paragraph by paragraph.

    Mirrors the body-paragraph path: compound markers first, then only an
    unambiguous marker/slot mapping. Anything ambiguous is left untouched so
    the template never gains a misassigned variable.
    """
    pending = list(slots)
    for paragraph in cell.paragraphs:
        if not pending:
            break
        text = paragraph.text or ''
        compound_keys = _replace_compound_visuals_in_paragraph(
            paragraph, text, pending,
        )
        pending = [slot for slot in pending if slot.key not in compound_keys]
        if not pending:
            break
        text = paragraph.text or ''
        present = [
            slot for slot in pending
            if any(raw in text for raw in slot.raw_markers)
        ]
        matches = list(_visual_matches(text))
        if not present or len(matches) != len(present):
            continue
        _replace_visuals_in_paragraph(paragraph, text, matches, present)
        present_keys = {slot.key for slot in present}
        pending = [slot for slot in pending if slot.key not in present_keys]


def _replace_compound_visuals_in_paragraph(
    paragraph: Any,
    text: str,
    slots_for_paragraph: list[Slot],
) -> set[str]:
    """Replace fields whose value owns surrounding grammar, not one fill line."""
    replacements = []
    replaced_keys = set()
    for slot in slots_for_paragraph:
        for raw in slot.raw_markers:
            if len(list(_visual_matches(raw))) < 2:
                continue
            start = text.find(raw)
            if start < 0:
                continue
            replacements.append((
                start,
                start + len(raw),
                f'{{{{ {slot.key} or {raw!r} }}}}',
            ))
            replaced_keys.add(slot.key)
            break
    _replace_text_spans_in_paragraph(paragraph, text, replacements)
    return replaced_keys


def _insert_slot_into_paragraph(paragraph: Any, key: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = f'{{{{ {key} }}}}'
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(f'{{{{ {key} }}}}')


def _set_cell_text(cell: Any, key: str, *, checkbox: bool = False) -> None:
    paragraphs = cell.paragraphs
    paragraph = paragraphs[0] if paragraphs else cell.add_paragraph()
    text = f'{{{{ {key} }}}}'
    if checkbox:
        text = '☐ ' + text
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(text)


def _location_parts(location: str) -> tuple[int, int, int] | None:
    """Parse ``body.tables[T].rows[R].cells[C]`` into a triple, else None."""
    match = re.fullmatch(r'body\.tables\[(\d+)\]\.rows\[(\d+)\]\.cells\[(\d+)\]', location)
    if match is None:
        return None
    return int(match.group(1)), int(match.group(2)), int(match.group(3))


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """Replace paragraph text while retaining the formatting of its first run."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(text)


def _normalize_numbered_paragraph_group(document: Any, slot: Slot) -> None:
    indices = sorted(
        int(location[len('body.paragraphs['):-1])
        for location in slot.locations
        if location.startswith('body.paragraphs[')
    )
    if len(indices) < 2 or slot.repeat is None:
        return

    paragraphs = document.paragraphs
    start = paragraphs[indices[0]]
    body = paragraphs[indices[1]]
    _set_paragraph_text(start, f'{{%p for item in {slot.key} %}}')
    _set_paragraph_text(
        body,
        f"{slot.repeat['major']}.{{{{ loop.index }}}}. {{{{ item or '___________________' }}}}",
    )

    if len(indices) >= 3:
        end = paragraphs[indices[2]]
        for index in reversed(indices[3:]):
            element = document.paragraphs[index]._element
            element.getparent().remove(element)
    else:
        end_element = deepcopy(body._element)
        body._element.addnext(end_element)
        end = Paragraph(end_element, body._parent)
    _set_paragraph_text(end, '{%p endfor %}')


def _normalize_outline_group(document: Any, slot: Slot) -> None:
    indices = [int(re.search(r'\[(\d+)\]', location).group(1)) for location in slot.locations]
    paragraphs = document.paragraphs
    prototypes = {}
    for index, level in zip(indices, slot.repeat['prototype_levels']):
        prototypes.setdefault(level, paragraphs[index])
    anchor = paragraphs[indices[0]]._p

    def insert(prototype, text, *, body=False):
        element = deepcopy(prototype._p)
        paragraph = Paragraph(element, prototype._parent)
        _set_paragraph_text(paragraph, text)
        numbering = _paragraph_numbering(prototype)
        if numbering is not None:
            indentation = numbering[1].findall(f"{qn('w:pPr')}/{qn('w:ind')}")
            properties = paragraph._p.get_or_add_pPr()
            if indentation and properties.ind is None:
                properties.append(deepcopy(indentation[0]))
        if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
            paragraph._p.pPr.remove(paragraph._p.pPr.numPr)
        # Override any numbering inherited from a paragraph style as well.
        if _paragraph_numbering(paragraph) is not None:
            paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0
        if body and '{{ row.number }}' in text:
            paragraph.paragraph_format.keep_with_next = len(prototypes) > 1 and prototype is prototypes[0]
        anchor.addprevious(element)

    insert(prototypes[0], f'{{%p for row in {slot.key} %}}')
    for level, prototype in sorted(prototypes.items()):
        insert(prototype, f"{{%p if row.level == '{level}' %}}")
        insert(prototype, '{{ row.number }}. {{ row.text }}', body=True)
        insert(prototype, '{%p endif %}')
    insert(prototypes[0], '{%p endfor %}')
    for index in indices:
        element = paragraphs[index]._p
        element.getparent().remove(element)


def _set_cell_content(cell: Any, text: str) -> None:
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    _set_paragraph_text(paragraph, text)
    for extra in cell.paragraphs[1:]:
        _set_paragraph_text(extra, '')


def _set_row_control(row: _Row, text: str) -> None:
    for cell_index, cell in enumerate(row.cells):
        _set_cell_content(cell, text if cell_index == 0 else '')


def _normalize_numbered_table_group(document: Any, slot: Slot) -> None:
    if slot.repeat is None:
        return
    table = document.tables[slot.repeat['table']]
    indices = sorted(
        int(location.rsplit('[', 1)[1][:-1])
        for location in slot.locations
        if location.startswith(f"body.tables[{slot.repeat['table']}].rows[")
    )
    if len(indices) < 2:
        return

    rows = table.rows
    start = rows[indices[0]]
    body = rows[indices[1]]
    _set_row_control(start, f'{{%tr for row in {slot.key} %}}')
    _set_cell_content(body.cells[0], '{{ loop.index }}.')
    for column in slot.repeat['columns']:
        _set_cell_content(body.cells[column['cell']], f"{{{{ row.{column['key']} or '___________________' }}}}")

    if len(indices) >= 3:
        end = rows[indices[2]]
        for index in reversed(indices[3:]):
            element = table.rows[index]._element
            element.getparent().remove(element)
    else:
        end_element = deepcopy(body._element)
        body._element.addnext(end_element)
        end = _Row(end_element, body._parent)
    _set_row_control(end, '{%tr endfor %}')


def normalize_visual_markers(docx_bytes: bytes) -> bytes:
    """Return a renderable copy where blanks and empty fields become variables.

    Visual ``___`` markers, ``[field description]`` instructions, grey-bordered
    empty paragraphs and empty table cells each get a named ``{{ key }}`` in the
    working copy. Visual markers
    remain as the fallback when a generated value is empty, preserving the
    source document's fillable spaces. The upload flow stores the source bytes
    separately, so this never alters the user's file.
    """
    parsed = parse_docx_template(docx_bytes, infer_structural=True)
    if parsed.render_ready:
        return docx_bytes

    visual_by_paragraph: dict[int, list[Slot]] = {}
    structural_paragraphs: dict[int, Slot] = {}
    table_cells: dict[tuple[int, int, int], Slot] = {}
    visual_cells: dict[tuple[int, int, int], list[Slot]] = {}
    checkbox_cells: dict[tuple[int, int], Slot] = {}
    repeated_paragraphs: list[Slot] = []
    repeated_tables: list[Slot] = []

    for slot in parsed.slots:
        if slot.repeat:
            if slot.repeat.get('kind') in ('numbered_paragraphs', 'numbered_outline'):
                repeated_paragraphs.append(slot)
            elif slot.repeat.get('kind') == 'table_rows':
                repeated_tables.append(slot)
            continue
        for location in slot.locations:
            if slot.source == 'visual' and location.startswith('body.paragraphs['):
                index = int(location[len('body.paragraphs['):-1])
                visual_by_paragraph.setdefault(index, []).append(slot)
            elif slot.source == 'visual' and location.startswith('body.tables['):
                parts = _location_parts(location)
                if parts is None:
                    continue
                visual_cells.setdefault(parts, []).append(slot)
            elif slot.source == 'structural' and location.startswith('body.paragraphs['):
                index = int(location[len('body.paragraphs['):-1])
                structural_paragraphs.setdefault(index, slot)
            elif slot.source == 'structural' and location.startswith('body.tables['):
                parts = _location_parts(location)
                if parts is None:
                    continue
                table_index, row_index, cell_index = parts
                if slot.raw_markers and slot.raw_markers[0].startswith('checkbox-choice:'):
                    checkbox_cells[(table_index, row_index)] = slot
                else:
                    table_cells[(table_index, row_index, cell_index)] = slot

    document = Document(BytesIO(docx_bytes))

    for index, slots in visual_by_paragraph.items():
        paragraph = document.paragraphs[index]
        text = paragraph.text or ''
        compound_keys = _replace_compound_visuals_in_paragraph(
            paragraph, text, slots,
        )
        simple_slots = [slot for slot in slots if slot.key not in compound_keys]
        text = paragraph.text or ''
        matches = list(_visual_matches(text))
        if len(matches) != len(simple_slots):
            continue  # ambiguous marker/slot mapping — leave the paragraph alone
        _replace_visuals_in_paragraph(paragraph, text, matches, simple_slots)

    for index, slot in structural_paragraphs.items():
        _insert_slot_into_paragraph(document.paragraphs[index], slot.key)

    for table_index, table in enumerate(document.tables):
        for (cell_table, row_index, cell_index), slot in table_cells.items():
            if cell_table == table_index:
                _set_cell_text(table.rows[row_index].cells[cell_index], slot.key)
        for (choice_table, row_index), slot in checkbox_cells.items():
            if choice_table == table_index:
                _set_cell_text(
                    table.rows[row_index].cells[0], slot.key, checkbox=True,
                )
        for (cell_table, row_index, cell_index), slots in visual_cells.items():
            if cell_table != table_index:
                continue
            _replace_visuals_in_cell(
                table.rows[row_index].cells[cell_index], slots,
            )

    for slot in sorted(
        repeated_tables,
        key=lambda candidate: (
            candidate.repeat['table'], candidate.repeat['start_row']
        ),
        reverse=True,
    ):
        _normalize_numbered_table_group(document, slot)

    for slot in sorted(
        repeated_paragraphs,
        key=lambda candidate: int(
            candidate.locations[0][len('body.paragraphs['):-1]
        ),
        reverse=True,
    ):
        if slot.repeat['kind'] == 'numbered_outline':
            _normalize_outline_group(document, slot)
        else:
            _normalize_numbered_paragraph_group(document, slot)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def prepare_flexible_template(
    docx_bytes: bytes, stored_slots: list[dict],
) -> tuple[bytes, ParsedTemplate]:
    """Upgrade old fixed outline fields in memory before a new generation."""
    source = parse_docx_template(docx_bytes)
    working = normalize_visual_markers(docx_bytes) if not source.render_ready else docx_bytes
    parsed = parse_docx_template(working) if working != docx_bytes else source
    metadata = {slot.key: asdict(slot) for slot in source.slots}
    metadata.update({slot['key']: slot for slot in stored_slots})
    # Newly detected structure takes precedence over an obsolete scalar contract.
    metadata.update({slot.key: asdict(slot) for slot in source.slots if slot.repeat})
    for slot in parsed.slots:
        stored = metadata.get(slot.key)
        if stored:
            slot.label = stored['label']
            slot.value_type = stored['value_type']
            slot.repeat = stored.get('repeat')
            slot.omit_when_empty = stored.get('omit_when_empty', False)
            slot.inline = stored.get('inline', False)
            slot.marker_contexts = stored.get('marker_contexts', [])
    return working, parsed


def parsed_to_descriptor(parsed: ParsedTemplate, name: str = "Template") -> dict[str, Any]:
    """Convert parsed metadata into the stored descriptor and JSON Schema."""
    properties: dict[str, Any] = {}
    for slot in parsed.slots:
        if slot.value_type == "list[string]":
            properties[slot.key] = {
                "type": "array",
                "title": slot.label,
                "items": {"type": "string"},
            }
        elif slot.value_type == 'list[object]' and slot.repeat:
            columns = slot.repeat.get('columns', [])
            properties[slot.key] = {
                'type': 'array',
                'title': slot.label,
                'items': {
                    'type': 'object',
                    'properties': {
                        column['key']: {
                            'type': 'string',
                            'title': column['label'],
                            **({'enum': [str(i) for i in range(slot.repeat['levels'])]}
                               if slot.repeat.get('kind') == 'numbered_outline'
                               and column['key'] == 'level' else {}),
                        }
                        for column in columns
                    },
                    'required': [column['key'] for column in columns],
                    'additionalProperties': False,
                },
            }
        else:
            properties[slot.key] = {"type": "string", "title": slot.label}

    return {
        "name": name,
        "slots": [asdict(slot) for slot in parsed.slots],
        "schema_json": {
            "type": "object",
            "properties": properties,
            "required": [slot.key for slot in parsed.slots if slot.required],
            "additionalProperties": False,
        },
        "style_config": parsed.style_config,
        "render_ready": parsed.render_ready,
        "warnings": parsed.warnings,
        "stats": {
            "paragraphs": len(parsed.paragraphs),
            "tables": len(parsed.tables),
            "headers": len(parsed.headers),
            "footers": len(parsed.footers),
            "slots": len(parsed.slots),
        },
    }


def validate_template_values(slots: list[Slot], values: dict[str, Any]) -> None:
    """Validate renderable structure, not content completeness or row counts."""
    if not isinstance(values, dict):
        raise ValueError('Template values must be a JSON object')
    expected = {slot.key for slot in slots}
    missing = sorted(expected - values.keys())
    if missing:
        raise ValueError(f"Template values are missing: {', '.join(missing)}")
    extra = sorted(values.keys() - expected)
    if extra:
        raise ValueError(f"Template values contain unknown fields: {', '.join(extra)}")

    for slot in slots:
        value = values[slot.key]
        if slot.value_type == "list[string]":
            if not isinstance(value, list) or not all(isinstance(item, str) for item in value):
                raise ValueError(f"Template value '{slot.key}' must be a list of strings")
        elif slot.value_type == 'list[object]':
            columns = slot.repeat.get('columns', []) if slot.repeat else []
            expected_columns = {column['key'] for column in columns}
            valid = isinstance(value, list) and all(
                isinstance(item, dict)
                and set(item) == expected_columns
                and all(isinstance(item[key], str) for key in expected_columns)
                for item in value
            )
            if not valid:
                raise ValueError(
                    f"Template value '{slot.key}' must be a list of row objects"
                )
            if slot.repeat and slot.repeat.get('kind') == 'numbered_outline':
                previous = -1
                for item in value:
                    level = item['level']
                    if level not in {str(i) for i in range(slot.repeat['levels'])}:
                        raise ValueError(f"Invalid outline level in '{slot.key}'")
                    if int(level) > previous + 1:
                        raise ValueError(f"Outline '{slot.key}' needs a parent")
                    previous = int(level)
        elif not isinstance(value, str):
            raise ValueError(f"Template value '{slot.key}' must be a string")


def _loop_iterables(docx_bytes: bytes) -> set[str]:
    """Slot keys used as ``{% for ... in <key> %}`` iterables.

    Such slots must keep their list value so the loop can iterate it; the
    plain ``{{ key }}`` case needs the list joined into one string instead.
    """
    with zipfile.ZipFile(BytesIO(docx_bytes)) as archive:
        try:
            xml = archive.read('word/document.xml').decode('utf-8', 'replace')
        except KeyError:
            return set()

    iterables: set[str] = set()
    for match in re.finditer(
        r'\{%[-+]?\s*(?:p|tr|tc|r)?\s*for\s+\w+\s+in\s+([^%}]+?)\s*[-+]?%\}',
        xml,
    ):
        name = re.match(r'[A-Za-z_][A-Za-z0-9_]*', match.group(1).strip())
        if name:
            iterables.add(name.group(0))
    return iterables


def _without_empty_omitted_paragraphs(
    docx_bytes: bytes,
    slots: list[Slot],
    values: dict[str, Any],
) -> bytes:
    omitted_indices = []
    for slot in slots:
        value = values.get(slot.key)
        if not slot.omit_when_empty or value not in ('', []):
            continue
        for location in slot.locations:
            match = re.fullmatch(r'body\.paragraphs\[(\d+)\]', location)
            if match:
                omitted_indices.append(int(match.group(1)))
    if not omitted_indices:
        return docx_bytes

    document = Document(BytesIO(docx_bytes))
    for index in sorted(set(omitted_indices), reverse=True):
        element = document.paragraphs[index]._element
        element.getparent().remove(element)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_docx_template(
    docx_bytes: bytes,
    values: dict[str, Any],
    *,
    slots: list[Slot] | None = None,
) -> bytes:
    """Fill a render-ready DOCX using its semantically enriched slot descriptor."""
    parsed = parse_docx_template(docx_bytes)
    if not parsed.render_ready:
        raise ValueError("Template has unnamed visual fields and must be reviewed first")

    descriptor_slots = parsed.slots if slots is None else slots
    parsed_keys = {slot.key for slot in parsed.slots}
    descriptor_keys = {slot.key for slot in descriptor_slots}
    if descriptor_keys != parsed_keys:
        raise ValueError("Template descriptor fields do not match the DOCX placeholders")

    validate_template_values(descriptor_slots, values)
    render_docx_bytes = _without_empty_omitted_paragraphs(
        docx_bytes, descriptor_slots, values,
    )
    descriptor_by_key = {slot.key: slot for slot in descriptor_slots}
    value_types = {slot.key: slot.value_type for slot in descriptor_slots}
    loop_iterables = _loop_iterables(render_docx_bytes)
    render_values = {
        slot.key: (
            values[slot.key]
            if slot.key in loop_iterables
            else ', '.join(values[slot.key])
            if value_types[slot.key] == 'list[string]'
            and descriptor_by_key[slot.key].inline
            else '\n'.join(values[slot.key])
            if value_types[slot.key] == 'list[string]'
            else values[slot.key]
        )
        for slot in parsed.slots
    }
    for slot in descriptor_slots:
        if slot.repeat and slot.repeat.get('kind') == 'numbered_outline':
            counters = [0] * slot.repeat['levels']
            rows = []
            for item in values[slot.key]:
                level = int(item['level'])
                counters[level] += 1
                counters[level + 1:] = [0] * (len(counters) - level - 1)
                rows.append({**item, 'number': '.'.join(map(str, counters[:level + 1]))})
            render_values[slot.key] = rows
    template = DocxTemplate(BytesIO(render_docx_bytes))
    template.render(
        render_values,
        jinja_env=SandboxedEnvironment(undefined=StrictUndefined),
    )
    output = BytesIO()
    template.save(output)
    return output.getvalue()


def render_empty_template_preview(docx_bytes: bytes) -> bytes:
    """Render the fillable DOCX with empty values for human-facing previews.

    The working template contains Jinja syntax required by docxtpl. It is never
    suitable for a browser preview: Word/document viewers show the syntax as
    ordinary text. This produces a separate display-only document and never
    changes the source upload or the template used for actual generation.
    """
    parsed = parse_docx_template(docx_bytes)
    if not parsed.render_ready:
        raise ValueError('Template preview requires render-ready fields')
    values = {
        slot.key: [] if slot.value_type.startswith('list[') else ''
        for slot in parsed.slots
    }
    return render_docx_template(docx_bytes, values, slots=parsed.slots)


def suggest_additional_prompt(parsed: ParsedTemplate) -> str:
    """Create a deterministic first draft; an admin may edit it before approval."""
    keys = ", ".join(slot.key for slot in parsed.slots)
    if not keys:
        return "Если сведений нет в транскрипте, не придумывай их."
    return (
        f"Заполни поля шаблона: {keys}. "
        "Используй только сведения из транскрипта и метаданных встречи. "
        "Если сведений для поля нет, верни пустую строку или массив; не галлюцинируй."
    )
