"""Owned generation lifecycle through real ASGI routes and isolated SQLite.

No user's database/settings, model services or installed CLI are accessed.
"""
import asyncio
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import AsyncMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
_storage = tempfile.TemporaryDirectory(prefix='steppe-cancel-import-')
with patch.dict(os.environ, STEPPE_DATA_DIR=_storage.name):
    import desktop_server as server
import httpx
from generation_jobs import GenerationJobs


class GenerationCancellationTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.storage = tempfile.TemporaryDirectory(prefix='steppe-cancel-test-')
        self.path = Path(self.storage.name)
        self.original_index = server.indexer.index_meeting
        self.patches = [patch.object(server, 'DB_PATH', self.path / 'test.db'),
                        patch.object(server, 'generation_jobs', GenerationJobs(server.write_generation_status)),
                        patch.object(server, 'sync_generation_settings'),
                        patch.object(server.indexer, 'index_meeting', AsyncMock(return_value=1))]
        for p in self.patches: p.start()
        server.init_db()
        server.RecordingStreams(server.get_db, server.get_setting)
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=server.app), base_url='http://test')
        self.id = await self.create()
        self.stage_patches = [patch.object(server.protocol, 'generate', AsyncMock(return_value={'title': 'new protocol'})),
                              patch.object(server.protocol, '_translate', AsyncMock(return_value={'title': 'new kz'})),
                              patch.object(server.summary, 'generate', AsyncMock(return_value={'overview': 'new summary'})),
                              patch.object(server.summary, 'translate', AsyncMock(return_value={'overview': 'new kz'}))]
        for p in self.stage_patches: p.start()

    async def asyncTearDown(self):
        await server.generation_jobs.shutdown()
        await self.client.aclose()
        for p in reversed(self.stage_patches + self.patches): p.stop()
        self.storage.cleanup()

    async def create(self):
        r = await self.client.post('/api/v1/meetings/create', json={'title': 'Synthetic', 'transcript_text': 'Синтетическая речь'})
        return r.json()['id']

    async def post(self, action, id=None, **kwargs):
        return await self.client.post(f'/api/v1/meetings/{id or self.id}/{action}', **kwargs)

    async def meeting(self, id=None):
        return (await self.client.get(f'/api/v1/meetings/{id or self.id}')).json()

    async def settle(self, id=None):
        job = server.generation_jobs.jobs.get(id or self.id)
        if job:
            await asyncio.wait_for(asyncio.shield(job.task), 3)
            await asyncio.sleep(0)

    def seed_prior(self):
        with server.get_db() as db:
            db.execute("UPDATE meetings SET protocol_ru=?, protocol_kz=?, summary_ru=?, summary_kz=?, template_values_ru=?, template_values_kz=?, transcript_segments=? WHERE id=?",
                       (*[json.dumps({'prior': column}) for column in ('ru', 'kz', 'sru', 'skz', 'tru', 'tkz')], '[{"text":"saved segment"}]', self.id))
            db.commit()

    async def test_stop_immediately_after_accept_is_owned_and_durable_then_explicit_rerun(self):
        self.seed_prior()
        prior = await self.meeting()
        async def pending_generation(**kwargs):
            await asyncio.Future()
        # ASGI middleware may yield before the response reaches the caller.
        # Keep inference genuinely running when testing Stop ownership.
        with patch.object(server.protocol, 'generate', pending_generation):
            response = await self.post('generate')
            self.assertEqual(response.status_code, 200)
            self.assertIsNotNone(server.generation_jobs.active(self.id))
            duplicate = await self.post('generate')
            self.assertEqual(duplicate.status_code, 409)
            stop = await self.post('stop-generation')
            self.assertEqual(stop.json()['status'], 'cancelled')
            self.assertEqual((await self.post('stop-generation')).json()['status'], 'cancelled')
            row = await self.meeting()
            self.assertEqual(row['status'], 'cancelled')
            for key in ('transcript_text', 'transcript_segments', 'protocol_ru', 'protocol_kz', 'summary_ru', 'summary_kz', 'template_values_ru', 'template_values_kz'):
                self.assertEqual(row[key], prior[key])
        await self.post('generate')
        await self.settle()
        self.assertEqual((await self.meeting())['status'], 'completed')
        self.assertEqual((await self.post('stop-generation')).json()['status'], 'completed')

    async def test_stop_waits_for_cleanup_duplicate_is_idempotent_and_other_meeting_unaffected(self):
        entered, cleanup, release = asyncio.Event(), asyncio.Event(), asyncio.Event()
        async def slow(**kwargs):
            entered.set()
            try:
                await asyncio.Future()
            finally:
                cleanup.set()
                await release.wait()
        with patch.object(server.protocol, 'generate', slow):
            await self.post('generate')
            await entered.wait()
            job = server.generation_jobs.jobs[self.id]
            first = asyncio.create_task(self.post('stop-generation'))
            await cleanup.wait()
            second = asyncio.create_task(self.post('stop-generation'))
            await asyncio.sleep(0)
            self.assertFalse(first.done())
            self.assertFalse(second.done())
            self.assertEqual(job.task.cancelling(), 1)
            self.assertEqual((await self.meeting())['status'], 'cancelling')
            self.assertEqual((await self.post('generate')).status_code, 409)
            other = await self.create()
            with patch.object(server.protocol, 'generate', AsyncMock(return_value={})):
                await self.post('generate', other)
                await self.settle(other)
            self.assertEqual((await self.meeting(other))['status'], 'completed')
            release.set()
            for result in await asyncio.gather(first, second):
                self.assertEqual(result.json()['status'], 'cancelled')
        self.assertTrue(job.task.cancelled())

    async def test_timeout_does_not_interrupt_cleanup_and_retry_acknowledges(self):
        entered, release = asyncio.Event(), asyncio.Event()
        async def slow(**kwargs):
            entered.set()
            try: await asyncio.Future()
            finally: await release.wait()
        with patch.object(server.protocol, 'generate', slow):
            await self.post('generate'); await entered.wait()
            with self.assertRaises(TimeoutError):
                await server.generation_jobs.stop(self.id, timeout=.01)
            self.assertEqual((await self.meeting())['status'], 'cancelling')
            release.set()
            result = await self.post('stop-generation')
            self.assertEqual(result.json()['status'], 'cancelled')

    async def test_swallowed_cancel_and_detached_old_write_cannot_overwrite_rerun(self):
        entered = asyncio.Event()
        delayed_write = None
        release = asyncio.Event()
        async def stubborn(**kwargs):
            nonlocal delayed_write
            async def stale():
                await release.wait()
                server.save_generation_results(self.id, {field: {'stale': True} for field in ('protocol_ru', 'protocol_kz', 'summary_ru', 'summary_kz', 'template_values_ru', 'template_values_kz')})
            delayed_write = asyncio.create_task(stale())  # inherits old job identity
            entered.set()
            try: await asyncio.Future()
            except asyncio.CancelledError: return {'stale': True}
        self.seed_prior()
        with patch.object(server.protocol, 'generate', stubborn):
            await self.post('generate'); await entered.wait()
            self.assertEqual((await self.post('stop-generation')).json()['status'], 'cancelled')
        self.assertEqual((await self.meeting())['protocol_ru'], {'prior': 'ru'})
        await self.post('generate'); await self.settle()
        release.set()
        with self.assertRaises(asyncio.CancelledError): await delayed_write
        self.assertEqual((await self.meeting())['protocol_ru'], {'title': 'new protocol'})
        self.assertEqual((await self.meeting())['status'], 'completed')

    async def test_template_stage_stays_generating_and_stop_preserves_entire_prior_set(self):
        from docx import Document
        template = self.path / 'template.docx'
        doc = Document(); doc.add_paragraph('{{ title }}'); doc.save(template)
        self.seed_prior()
        with server.get_db() as db:
            db.execute("INSERT INTO protocol_templates(id,name,docx_path,source_docx_path,slots,schema_json,style_config,created_at) VALUES ('t','t',?,?,'[]','{}','{}','now')", (str(template), str(template)))
            db.execute("UPDATE meetings SET template_id='t' WHERE id=?", (self.id,)); db.commit()
        entered = asyncio.Event()
        async def slow(*args, **kwargs):
            entered.set(); await asyncio.Future()
        with patch.object(server.template_generation, 'generate_protocol_template_values', slow):
            await self.post('generate'); await entered.wait()
            row = await self.meeting()
            self.assertEqual(row['status'], 'generating')
            self.assertEqual(row['protocol_ru'], {'prior': 'ru'})
            self.assertEqual(row['summary_ru'], {'prior': 'sru'})
            await self.post('stop-generation')
        row = await self.meeting()
        self.assertEqual(row['status'], 'cancelled')
        self.assertEqual(row['template_values_ru'], {'prior': 'tru'})
        self.assertEqual(row['template_values_kz'], {'prior': 'tkz'})
        server.indexer.index_meeting.assert_not_awaited()

    async def test_stop_at_each_translation_summary_stage_preserves_entire_prior_set(self):
        for module, method in ((server.protocol, '_translate'),
                               (server.summary, 'generate'), (server.summary, 'translate')):
            with self.subTest(stage=method):
                self.seed_prior()
                prior = await self.meeting()
                entered = asyncio.Event()
                async def slow(*args, **kwargs): entered.set(); await asyncio.Future()
                with patch.object(module, method, slow):
                    await self.post('generate'); await entered.wait()
                    await self.post('stop-generation')
                row = await self.meeting()
                for field in ('protocol_ru', 'protocol_kz', 'summary_ru', 'summary_kz', 'template_values_ru', 'template_values_kz'):
                    self.assertEqual(row[field], prior[field])
                self.assertEqual(row['status'], 'cancelled')

    async def test_stop_during_template_translation_preserves_entire_prior_set(self):
        from docx import Document
        template = self.path / 'translation.docx'
        doc = Document(); doc.add_paragraph('{{ title }}'); doc.save(template)
        self.seed_prior()
        with server.get_db() as db:
            db.execute("INSERT INTO protocol_templates(id,name,docx_path,source_docx_path,slots,schema_json,style_config,created_at) VALUES ('t','t',?,?,'[]','{}','{}','now')", (str(template), str(template)))
            db.execute("UPDATE meetings SET template_id='t' WHERE id=?", (self.id,)); db.commit()
        entered = asyncio.Event()
        async def stubborn(*args, **kwargs):
            entered.set()
            try: await asyncio.Future()
            except asyncio.CancelledError: return {'stale': 'translation'}
        with patch.object(server.template_generation, 'generate_protocol_template_values', AsyncMock(return_value={'template_values': {'title': 'new ru'}})), \
             patch.object(server.template_generation, 'translate_template_values', stubborn):
            await self.post('generate'); await entered.wait()
            await self.post('stop-generation')
        row = await self.meeting()
        self.assertEqual(row['template_values_ru'], {'prior': 'tru'})
        self.assertEqual(row['template_values_kz'], {'prior': 'tkz'})
        self.assertEqual(row['status'], 'cancelled')

    async def test_stage_error_preserves_all_prior_results(self):
        self.seed_prior()
        prior = await self.meeting()
        with patch.object(server.summary, 'translate', AsyncMock(side_effect=RuntimeError('Synthetic failure'))):
            await self.post('generate'); await self.settle()
        row = await self.meeting()
        self.assertEqual(row['status'], 'error')
        for field in ('protocol_ru', 'protocol_kz', 'summary_ru', 'summary_kz', 'template_values_ru', 'template_values_kz'):
            self.assertEqual(row[field], prior[field])

    async def test_template_failure_keeps_old_set_then_success_publishes_new_set_atomically(self):
        from docx import Document
        template = self.path / 'template-result.docx'
        doc = Document(); doc.add_paragraph('{{ title }}'); doc.save(template)
        self.seed_prior()
        prior = await self.meeting()
        with server.get_db() as db:
            db.execute("INSERT INTO protocol_templates(id,name,docx_path,source_docx_path,slots,schema_json,style_config,created_at) VALUES ('t','t',?,?,'[]','{}','{}','now')", (str(template), str(template)))
            db.execute("UPDATE meetings SET template_id='t' WHERE id=?", (self.id,)); db.commit()
        with patch.object(server.template_generation, 'generate_protocol_template_values', AsyncMock(return_value={'template_values': {'title': 'new ru'}})), \
             patch.object(server.template_generation, 'translate_template_values', AsyncMock(side_effect=RuntimeError('Synthetic template failure'))):
            await self.post('generate'); await self.settle()
        row = await self.meeting()
        self.assertEqual(row['status'], 'error')
        for field in ('protocol_ru', 'protocol_kz', 'summary_ru', 'summary_kz', 'template_values_ru', 'template_values_kz'):
            self.assertEqual(row[field], prior[field])
        with patch.object(server.template_generation, 'generate_protocol_template_values', AsyncMock(return_value={'template_values': {'title': 'new ru'}})), \
             patch.object(server.template_generation, 'translate_template_values', AsyncMock(return_value={'title': 'new kz'})):
            await self.post('generate'); await self.settle()
        row = await self.meeting()
        self.assertEqual(row['status'], 'completed')
        for field in ('protocol_ru', 'protocol_kz', 'summary_ru', 'summary_kz', 'template_values_ru', 'template_values_kz'):
            self.assertNotEqual(row[field], prior[field])
        self.assertEqual(row['template_values_ru'], {'title': 'new ru'})
        self.assertEqual(row['template_values_kz'], {'title': 'new kz'})

    async def test_cancel_cleanup_exception_is_terminal_and_retry_is_idempotent(self):
        entered = asyncio.Event()
        async def slow(**kwargs):
            entered.set()
            try: await asyncio.Future()
            finally: raise RuntimeError('Synthetic cleanup exception')
        with patch.object(server.protocol, 'generate', slow):
            await self.post('generate'); await entered.wait()
            self.assertEqual((await self.post('stop-generation')).json()['status'], 'cancelled')
            self.assertEqual((await self.post('stop-generation')).json()['status'], 'cancelled')
            self.assertIsNone(server.generation_jobs.active(self.id))

    async def test_restart_recovers_interrupted_transcription_for_explicit_retry(self):
        with server.get_db() as db:
            db.execute("UPDATE meetings SET status='transcribing' WHERE id=?", (self.id,)); db.commit()
        server.recover_orphaned_generations()
        self.assertEqual((await self.meeting())['status'], 'error')
        self.assertIn('перезапуском', (await self.meeting())['error_message'])
        await self.post('generate'); await self.settle()
        self.assertEqual((await self.meeting())['status'], 'completed')

    async def test_disconnected_stop_request_cannot_cancel_provider_cleanup(self):
        entered, cleanup, release = asyncio.Event(), asyncio.Event(), asyncio.Event()
        async def slow(**kwargs):
            entered.set()
            try: await asyncio.Future()
            finally: cleanup.set(); await release.wait()
        with patch.object(server.protocol, 'generate', slow):
            await self.post('generate'); await entered.wait()
            job = server.generation_jobs.jobs[self.id]
            request = asyncio.create_task(self.post('stop-generation'))
            await cleanup.wait()
            request.cancel()
            with self.assertRaises(asyncio.CancelledError): await request
            self.assertEqual(job.task.cancelling(), 1)
            self.assertFalse(job.task.done())
            release.set()
            self.assertEqual((await self.post('stop-generation')).json()['status'], 'cancelled')

    async def test_index_continuation_checks_job_after_embedding_even_if_cancel_swallowed(self):
        entered = asyncio.Event()
        async def stubborn(texts):
            entered.set()
            try: await asyncio.Future()
            except asyncio.CancelledError: return [[0.0] for _ in texts]
        from engine import indexer
        # Use the real indexer instead of the ordinary no-network fixture.
        original = self.original_index
        with patch.object(server.indexer, 'index_meeting', original), \
             patch.object(indexer.embeddings, 'embed_texts', stubborn), \
             patch.object(indexer.vector_store, 'upsert_chunks') as write:
            await self.post('generate'); await entered.wait()
            await self.post('stop-generation')
            write.assert_not_called()
        self.assertEqual((await self.meeting())['status'], 'completed')

    async def test_auto_generate_after_transcription_hands_off_to_owned_job(self):
        audio = self.path / 'audio.wav'; audio.write_bytes(b'synthetic audio')
        with server.get_db() as db:
            db.execute("UPDATE meetings SET audio_path=?, transcript_text='', status='draft' WHERE id=?", (str(audio), self.id)); db.commit()
        entered = asyncio.Event()
        async def slow(**kwargs): entered.set(); await asyncio.Future()
        with patch.object(server.whisper_service, 'transcribe_local', return_value={'text': 'Распознано', 'segments': [{'text': 'сегмент'}], 'duration': 2}), \
             patch.object(server, 'get_setting', side_effect=lambda key, default='': default), \
             patch.object(server.protocol, 'generate', slow):
            response = await self.post('transcribe', params={'auto_generate': 'true'})
            self.assertEqual(response.status_code, 200)
            await entered.wait()
            await self.post('stop-generation')
            # Mini's offline completion retry is not an explicit rerun.
            retry = await self.post('transcribe', params={'auto_generate': 'true'})
            self.assertEqual(retry.json()['status'], 'cancelled')
        row = await self.meeting()
        self.assertEqual(row['status'], 'cancelled')
        self.assertEqual(row['transcript_text'], 'Распознано')
        self.assertEqual(row['transcript_segments'], [{'text': 'сегмент'}])
        self.assertTrue(audio.exists())

    async def test_mini_finalized_manifest_does_not_resubmit_cancelled_generation(self):
        audio = self.path / 'audio.wav'; audio.write_bytes(b'audio')
        with server.get_db() as db:
            db.execute("UPDATE meetings SET audio_path=? WHERE id=?", (str(audio), self.id))
            db.execute("INSERT INTO recording_streams(id,meeting_id,config,finalized) VALUES ('s',?,'{}',0)", (self.id,)); db.commit()
        def prepare(db, meeting, *args):
            db.execute("UPDATE meetings SET status='generating' WHERE id=?", (meeting,))
            db.execute("UPDATE recording_streams SET finalized=1 WHERE id='s'")
        manifest = {'session_id': 's', 'chunk_count': 1, 'duration_seconds': 1}
        async def pending_generation(**kwargs):
            await asyncio.Future()
        with patch.object(server.recording_streams, 'prepare_completion', prepare), \
             patch.object(server.protocol, 'generate', pending_generation):
            response = await self.post('process-recording', json=manifest)
            self.assertEqual(response.status_code, 200)
            self.assertIsNotNone(server.generation_jobs.active(self.id))
            await self.post('stop-generation')
            retry = await self.post('process-recording', json=manifest)
            self.assertEqual(retry.json()['status'], 'cancelled')
            self.assertIsNone(server.generation_jobs.active(self.id))
        await self.post('generate'); await self.settle()
        self.assertEqual((await self.meeting())['status'], 'completed')

    async def test_shutdown_cleanup_and_restart_recovery_preserve_results(self):
        self.seed_prior()
        entered, cleaned = asyncio.Event(), asyncio.Event()
        async def slow(**kwargs):
            entered.set()
            try: await asyncio.Future()
            finally: cleaned.set()
        with patch.object(server.protocol, 'generate', slow):
            await self.post('generate'); await entered.wait()
            await server.generation_jobs.shutdown()
        self.assertTrue(cleaned.is_set())
        self.assertEqual((await self.meeting())['status'], 'cancelled')
        with server.get_db() as db:
            db.execute("UPDATE meetings SET status='generating' WHERE id=?", (self.id,)); db.commit()
        server.recover_orphaned_generations()
        self.assertEqual((await self.meeting())['status'], 'cancelled')
        self.assertEqual((await self.meeting())['protocol_ru'], {'prior': 'ru'})

    async def test_all_four_cli_providers_stop_the_owned_process_via_meeting_api(self):
        from engine import cli_llm, llm
        async def real_llm(**kwargs):
            return await llm.complete_json([{'role': 'user', 'content': 'Synthetic JSON'}])
        for provider in ('codex', 'claude', 'antigravity', 'opencode'):
            with self.subTest(provider=provider):
                capture = self.path / f'{provider}.jsonl'
                binary = self.path / provider
                if provider == 'opencode':
                    source = f'CAPTURE={str(capture)!r}\nEVENTS=[]\nMODE=""\n' + (Path(__file__).parent / 'fixtures/cli/opencode-server.py').read_text()
                    model = 'local/exact#precise'
                else:
                    source = ('import json, os, sys, time\nfrom pathlib import Path\n'
                              f'Path({str(capture)!r}).write_text(json.dumps({{"pid":os.getpid(), "argv":sys.argv[1:], "cwd":os.getcwd()}})+"\\n")\n'
                              'sys.stdin.read()\ntime.sleep(60)\n')
                    model = 'exact-model'
                binary.write_text(f'#!{Path(sys.executable).resolve()}\n' + source)
                binary.chmod(0o700)
                with patch.object(server.protocol, 'generate', real_llm), \
                     patch.object(llm.settings, 'LLM_MODE', 'cli'), \
                     patch.object(llm.settings, 'LLM_CLI_PROVIDER', provider), \
                     patch.object(llm.settings, 'LLM_CLI', {provider: {'path': str(binary), 'model': model}}):
                    await self.post('generate')
                    async with asyncio.timeout(5):
                        while True:
                            if capture.exists():
                                records = [json.loads(line) for line in capture.read_text().splitlines()]
                                if provider != 'opencode' or any(c.get('path', '').endswith('/prompt') for c in records): break
                            await asyncio.sleep(.01)
                    result = await self.post('stop-generation')
                    self.assertEqual(result.json()['status'], 'cancelled')
                records = [json.loads(line) for line in capture.read_text().splitlines()]
                start = records[0]
                with self.assertRaises(ProcessLookupError): os.kill(start['pid'], 0)
                self.assertFalse(Path(start['cwd']).exists())
                if provider == 'opencode':
                    create = next(c['body'] for c in records if c.get('path') == '/api/session')
                    self.assertNotIn('id', create)
                    self.assertEqual(create['model'], {'providerID': 'local', 'id': 'exact', 'variant': 'precise'})
                    owned = next(c['path'].split('/')[3] for c in records if c.get('path', '').endswith('/prompt'))
                    self.assertEqual([(c['method'], c['path']) for c in records[-2:]],
                                     [('POST', f'/api/session/{owned}/interrupt'), ('DELETE', f'/api/session/{owned}')])
                else:
                    self.assertIn(model, start['argv'])
                self.assertEqual((await self.meeting())['status'], 'cancelled')

    async def test_http_generation_cancels_and_awaits_its_own_request(self):
        from engine import llm
        entered, cleaned = asyncio.Event(), asyncio.Event()
        async def post(*args, **kwargs):
            entered.set()
            try: await asyncio.Future()
            finally: cleaned.set()
        async def real_llm(**kwargs):
            return await llm.complete_json([{'role': 'user', 'content': 'Synthetic JSON'}])
        with patch.object(server.protocol, 'generate', real_llm), \
             patch.object(llm.settings, 'LLM_MODE', 'api'), \
             patch.object(llm, 'client', return_value=type('Client', (), {'post': staticmethod(post)})()), \
             patch.object(llm, '_save_call_log', AsyncMock()), \
             patch.object(llm.telemetry, 'emit', AsyncMock()):
            await self.post('generate'); await entered.wait()
            self.assertEqual((await self.post('stop-generation')).json()['status'], 'cancelled')
            self.assertTrue(cleaned.is_set())

    async def test_stop_missing_and_non_generating_meetings(self):
        self.assertEqual((await self.post('stop-generation', 'missing')).status_code, 404)
        self.assertEqual((await self.post('stop-generation')).json()['status'], 'transcribed')
